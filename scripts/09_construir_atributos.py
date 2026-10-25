"""Atributos de referencia: rutas completas entre anclajes explícitos de ARISTO.

Sin expansión de demanda, SBO, utilidades, coeficientes ni probabilidades logit.
"""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from zipfile import ZipFile
import calendar
import hashlib
import json
import math
import sqlite3
import sys

import geopandas as gpd
import pandas as pd
from openpyxl import load_workbook
from pyproj import Geod

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
from demanda.atributos import calendario_semanal,grafo_vial,camino_minimo,comprobar_secuencia,tiempos_aditivos

OUT=ROOT/'outputs/atributos_alternativas'
BASE=ROOT.parent
NAN=math.nan


def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def save(d,n):d.to_csv(OUT/f'{n}.csv',index=False,encoding='utf-8-sig',na_rep='NaN')
def rnd(v):return int(Decimal(str(v)).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
def md(d):
    def fmt(x):
        if pd.isna(x):return 'Pendiente'
        if isinstance(x,float):return f'{x:,.3f}'
        return str(x).replace('|','/').replace('\n',' ')
    return '\n'.join(['| '+' | '.join(d.columns)+' |','| '+' | '.join(['---']*len(d.columns))+' |']+
                     ['| '+' | '.join(fmt(v) for v in r)+' |' for r in d.itertuples(index=False,name=None)])


def horarios(cfg):
    departures=[];intervals=[];metrics=[];lookup={}
    for name,c in cfg['calendarios'].items():
        ordinary=[(d,h) for d,h,danger in c['salidas'] if not danger]
        m=calendario_semanal(ordinary); target=8 if c['regimen']=='baja' else 11
        for day,hour,danger in c['salidas']:
            departures.append(dict(calendario=name,sentido=c['sentido'],regimen=c['regimen'],dia_lunes0=day,
                                   hora=hour,carga_peligrosa=danger,incluido_ordinario=not danger,
                                   fuente=cfg['fuente_calendarios'],vigencia_fuente=c['fuente_vigencia']))
        for i,(start,gap) in enumerate(zip(m['salidas_h_semana'],m['intervalos_h']),1):
            intervals.append(dict(calendario=name,sentido=c['sentido'],orden=i,salida_h_semana=start,
                                  siguiente_h_semana=start+gap,intervalo_h=gap,incluye_cierre_semanal=i==m['f']))
        mm={k:v for k,v in m.items() if k not in ['salidas_h_semana','intervalos_h']}
        metrics.append(dict(calendario=name,sentido=c['sentido'],regimen=c['regimen'],**mm,
                            f_modelo=target,coincide_frecuencia=m['f']==target,
                            t_horario_calendario_referencia_h=cfg['w_horario']*m['H_ponderado_tiempo_h'],
                            vigencia=c['fuente_vigencia']))
        lookup[c['regimen'],c['sentido']]=(name,m,c)
    monthly=[]
    for station,months in cfg['estaciones'].items():
        for month in months:
            f=cfg['frecuencia_ordinaria_por_mes'][str(month)]
            regime='alta' if f==11 else 'baja'
            for direction in ['Chile Chico → Puerto Ibáñez','Puerto Ibáñez → Chile Chico']:
                name,m,c=lookup[regime,direction]; matched=m['f']==f
                monthly.append(dict(estacion=station,mes=month,dias=calendar.monthrange(cfg['ano_calendario'],month)[1],
                                    sentido_cruce=direction,f_ordinario=f,f_calendario_publicado=m['f'],calendario_fuente=name,
                                    H_168_f_h=168/f,H_media_intervalos_fuente_h=m['H_media_intervalos_h'],
                                    H_calendario_h=m['H_ponderado_tiempo_h'] if matched else NAN,
                                    intervalo_min_fuente_h=m['intervalo_min_h'],intervalo_max_fuente_h=m['intervalo_max_h'],
                                    t_horario_168_f_h=cfg['w_horario']*168/f,
                                    t_horario_calendario_h=cfg['w_horario']*m['H_ponderado_tiempo_h'] if matched else NAN,
                                    estado='calendario_referencia2026_transferencia_historica_condicionada' if matched else 'no_se_identifica_calendario11; fuente_tiene12',
                                    fuente=cfg['fuente_calendarios']))
    monthly=pd.DataFrame(monthly);seasons=[]
    for (station,direction),g in monthly.groupby(['estacion','sentido_cruce'],sort=False):
        full=g.H_calendario_h.notna().all();known=g.H_calendario_h.notna();days=g.dias.sum()
        seasons.append(dict(estacion=station,sentido_cruce=direction,
                            regimenes_mensuales=';'.join(f'{r.mes}:{r.f_ordinario}' for r in g.itertuples()),
                            H_168_f_h=(g.H_168_f_h*g.dias).sum()/days,
                            H_calendario_h=(g.H_calendario_h*g.dias).sum()/days if full else NAN,
                            t_horario_168_f_h=(g.t_horario_168_f_h*g.dias).sum()/days,
                            t_horario_calendario_h=(g.t_horario_calendario_h*g.dias).sum()/days if full else NAN,
                            dias_con_calendario_coincidente=int(g.loc[known,'dias'].sum()),dias_estacion=int(days),
                            calendario_estacion_completo=bool(full),
                            H_calendario_parcial_meses_conocidos_h=(g.loc[known,'H_calendario_h']*g.loc[known,'dias']).sum()/g.loc[known,'dias'].sum() if known.any() else NAN))
    return pd.DataFrame(departures),pd.DataFrame(intervals),pd.DataFrame(metrics),monthly,pd.DataFrame(seasons)


def cruce(frame,a,b):
    x=frame.loc[(frame.Nodo_A.eq(a)&frame.Nodo_B.eq(b))|(frame.Nodo_A.eq(b)&frame.Nodo_B.eq(a))]
    if len(x)!=1:raise ValueError(f'Cruce ambiguo/ausente: {a},{b}')
    return dict(x.iloc[0],desde=a,hasta=b)


def construir_ruta(frame,graphs,cfg,origin,destination,seg,alt,entry=None):
    o=cfg['puntos_referencia'][origin]['nodo_fisico'];d=cfg['puntos_referencia'][destination]['nodo_fisico']
    gch,gar=graphs['CAM' if seg=='CAM' else 'VL']
    if alt=='terrestre_chile':edges=camino_minimo(gch,o,d)
    elif alt=='barcaza_actual':
        forward=origin=='Chile Chico';a=cfg['barcaza']['Chile Chico' if forward else 'Puerto Ibáñez'];b=cfg['barcaza']['Puerto Ibáñez' if forward else 'Chile Chico']
        edges=camino_minimo(gch,o,a)+[cruce(frame,a,b)]+camino_minimo(gch,b,d)
    else:
        j=cfg['fronteras']['Jeinimeni'];e=cfg['fronteras'][entry]
        first,last=(j,e) if origin=='Chile Chico' else (e,j)
        edges=(camino_minimo(gch,o,first['Chile'])+[cruce(frame,first['Chile'],first['Argentina'])]+
               camino_minimo(gar,first['Argentina'],last['Argentina'])+[cruce(frame,last['Argentina'],last['Chile'])]+
               camino_minimo(gch,last['Chile'],d))
    comprobar_secuencia(edges,o,d)
    return edges


def medir_ruta(edges,seg,cfg):
    road=[e for e in edges if e['Clave']!='TRANSB'];boat=[e for e in edges if e['Clave']=='TRANSB']
    col='VEL VP' if seg=='CAM' else 'VEL VL';cost='CO VP' if seg=='CAM' else 'CO VL'
    drive=sum(e['LONGITUD']/e[col] for e in road)
    return dict(distancia_terrestre_total_km=sum(e['LONGITUD'] for e in road),
                tiempo_conduccion_red_h=drive,tiempo_navegacion_h=cfg['barcaza']['navegacion_h'] if boat else 0.0,
                tiempo_navegacion_longitud_velocidad_fuente_h=sum(e['LONGITUD']/e[col] for e in boat),
                distancia_navegacion_km=sum(e['LONGITUD'] for e in boat),
                COV_terrestre_ARISTO_CLP=sum(e[cost] for e in road),
                distancia_acceso_rampas_km=sum(e['LONGITUD'] for e in road if e['RUTA']=='RAMPA'),
                tiempo_rampas_red_h=sum(e['LONGITUD']/e[col] for e in road if e['RUTA']=='RAMPA'))


def auditar_costos(frame,used,costpath):
    f=load_workbook(costpath,data_only=False,keep_links=False);v=load_workbook(costpath,data_only=True,keep_links=False)
    sf=f['Costos_Red_Vial'];sv=v['Costos_Red_Vial'];tf=f['Tipos de arcos'];tv=v['Tipos de arcos']
    keyrow={tv.cell(i,9).value:i for i in range(2,55) if tv.cell(i,9).value is not None}
    formulas=[];checks=[]
    used_rows=frame.loc[frame.ID_Arco.isin(used)]
    for idx,r in used_rows.iterrows():
        row=int(idx)+2
        for c in ['I','J','Q','R','S']:
            formulas.append(dict(archivo='Cap 7/HDM4/Costos_Red_Vial_Actual.xlsx',hoja='Costos_Red_Vial',celda=f'{c}{row}',formula=sf[f'{c}{row}'].value,valor_cache=sv[f'{c}{row}'].value))
        if r.Clave in ['C','TRANSB']:continue
        k=keyrow[r.Clave];ratevl=tv.cell(k,18).value;ratevp=tv.cell(k,19).value
        qvl=rnd(Decimal(str(ratevl))*Decimal(str(r.LONGITUD)));qvp=rnd(Decimal(str(ratevp))*Decimal(str(r.LONGITUD)))
        cg=rnd(float(qvl)+(1 if r.TIPO_PAVIM=='P' else 1.5)*r.LONGITUD/r['VEL VL']*tv['V15'].value)
        checks.append(dict(ID_Arco=int(r.ID_Arco),fila_excel=row,km=r.LONGITUD,clave=r.Clave,
                           COV_VL_fuente=r['CO VL'],COV_VL_reconstruido=qvl,COV_CAM_fuente=r['CO VP'],COV_CAM_reconstruido=qvp,
                           CostGen_VL_fuente=r['CostGen VL'],CostGen_VL_reconstruido=cg,
                           diferencia_CO_VL=qvl-r['CO VL'],diferencia_CO_CAM=qvp-r['CO VP'],diferencia_CG=cg-r['CostGen VL']))
    for k in sorted({keyrow[r.Clave] for r in used_rows.itertuples() if r.Clave in keyrow}):
        for c in ['I','P','Q','R','S']:
            formulas.append(dict(archivo='Cap 7/HDM4/Costos_Red_Vial_Actual.xlsx',hoja='Tipos de arcos',celda=f'{c}{k}',formula=tf[f'{c}{k}'].value,valor_cache=tv[f'{c}{k}'].value))
    for c in ['V2','W2','V3','W3','V15','V44','U45','V45','Z71','AA71','AB71']:
        formulas.append(dict(archivo='Cap 7/HDM4/Costos_Red_Vial_Actual.xlsx',hoja='Tipos de arcos',celda=c,formula=tf[c].value,valor_cache=tv[c].value))
    params=dict(fecha_precios=str(tv['V44'].value),peso_CA=float(tv['V3'].value),peso_CS=float(tv['W3'].value),VST_CG_fuente_no_usado=float(tv['V15'].value),tarifa_VL_embebida_red=float(tv['Z71'].value),tarifa_CAM_embebida_red=float(tv['AA71'].value))
    f.close();v.close()
    return pd.DataFrame(formulas),pd.DataFrame(checks),params


def fila_atributos(origin,destination,seg,station,alt,rid,metrics,borders,schedule,cfg):
    boat=alt=='barcaza_actual';arg=alt=='terrestre_argentina';vehicle=seg in ['VL','CAM']
    roadtime=metrics['tiempo_conduccion_red_h'] if vehicle else NAN
    cov=metrics['COV_terrestre_ARISTO_CLP'] if vehicle else NAN
    nav=metrics['tiempo_navegacion_h'];anticip=cfg['anticipacion_h'][seg] if boat else 0.0
    fare=(cfg['tarifas_ARISTO_tabla52_p121']['VL_hasta_5_5m'] if seg=='VL' else
          cfg['tarifas_ARISTO_tabla52_p121']['adulto_general'] if seg=='PEATON' else NAN) if boat else 0.0
    identified=(cov if vehicle else 0.0)+(fare if math.isfinite(fare) else 0.0) if vehicle or boat else NAN
    m=tiempos_aditivos(roadtime,nav,boat)
    direction='Chile Chico → Puerto Ibáñez' if origin=='Chile Chico' else 'Puerto Ibáñez → Chile Chico'
    s=schedule.loc[schedule.estacion.eq(station)&schedule.sentido_cruce.eq(direction)].iloc[0]
    known_time=(roadtime if math.isfinite(roadtime) else 0)+nav+anticip if vehicle or boat else NAN
    pending=['localizacion_intrazonal_y_conectores_reales','peajes_y_otras_cargas_monetarias_no_verificadas']
    if not vehicle:pending+=['modo_y_tarifa_acceso_terrestre_persona','tiempo_acceso_persona; no_usar_velocidad_VL']
    if boat:
        pending+=['tiempo_desembarque','tasas_terminales','pasajeros_pagadores_por_vehiculo_y_condicion_tarifaria']
        if seg=='CAM':pending+=['longitud_vehiculo_carga_para_tarifa_lineal']
        if not s.calendario_estacion_completo:pending+=['calendario_ordinario11_zarpes_no_identificado']
    if arg:pending+=['tiempos_tramites_fronterizos','penalizacion_por_horarios_de_frontera','costos_privados_Argentina_y_moneda','seguro_y_permisos_internacionales']
    return dict(origen=origin,destino=destination,sentido=f'{origin} → {destination}',segmento=seg,estacion=station,alternativa=alt,ruta_id=rid,
                factibilidad_red=True,factibilidad_segmento=('condicionada_al_modo_terrestre_para_peaton' if not vehicle else 'condicionada_a_horarios_documentos_y_restricciones' if arg or boat else 'red_conexa_restricciones_peso_y_sentidos_urbanos_no_verificados'),
                alcance_OD=cfg['alcance_atributos'],origen_anclaje=cfg['puntos_referencia'][origin]['nodo_fisico'],destino_anclaje=cfg['puntos_referencia'][destination]['nodo_fisico'],
                unidad_usuario='vehiculo_liviano' if seg=='VL' else 'camion_clase_VP_modelada' if seg=='CAM' else 'persona_peaton_modo_acceso_no_identificado',
                distancia_terrestre_total_km=metrics['distancia_terrestre_total_km'],distancia_navegacion_km=metrics['distancia_navegacion_km'],
                tiempo_conduccion_terrestre_h=roadtime,tiempo_navegacion_h=nav,
                tiempo_acceso_terrestre_terminales_h=roadtime if boat else 0.0,
                tiempo_acceso_rampas_modelo_h=metrics['tiempo_rampas_red_h'] if vehicle else NAN,
                distancia_acceso_rampas_km=metrics['distancia_acceso_rampas_km'],
                **m,definicion_tiempos='barcaza: acceso=todo_tramo_terrestre_OD-terminal; viaje=navegacion. tierra: acceso0; viaje=conduccion. conduccion_descriptiva_no_se_suma_dos_veces',
                anticipacion_exigida_h=anticip,t_embarque_previo_h=anticip,operaciones_embarque_adicionales_h=0.0 if boat else 0.0,
                motivo_embarque_adicional0='incluidas_en_anticipacion_por_hipotesis_no_duplicar; no_duracion_medida' if boat else 'no_aplica',
                t_desembarque_h=NAN if boat else 0.0,t_embarque_desembarque_completo_h=NAN if boat else 0.0,
                H_h=s.H_calendario_h if boat else NAN,H_168_f_h=s.H_168_f_h if boat else NAN,
                t_horario_h=s.t_horario_calendario_h if boat else 0.0,t_horario_168_f_h=s.t_horario_168_f_h if boat else 0.0,w_horario=cfg['w_horario'],
                regimenes_frecuencia=s.regimenes_mensuales if boat else 'no_aplica',
                calendario_completo_estacion=bool(s.calendario_estacion_completo) if boat else True,
                fuente_H=cfg['fuente_calendarios'] if boat else 'sin_servicio_programado_en_red; frontera_separada',
                t_frontera_tramites_h=NAN if arg else 0.0,t_frontera_restriccion_horaria_h=NAN if arg else 0.0,
                tiempo_fisico_referencia_completo_h=roadtime if not boat and not arg and vehicle else NAN,
                subtotal_tiempos_identificados_h=known_time,
                subtotal_identificado_con_horario_simple_h=known_time+(s.t_horario_168_f_h if boat else 0),
                COV_terrestre_privado_ARISTO_CLP=cov,base_precios_COV='CLP_diciembre2025_TiposV44; precios_privados_V45',
                tarifa_transbordador_principal_CLP=fare,
                tarifa_CAM_CLP_por_metro=cfg['tarifas_ARISTO_tabla52_p121']['metro_lineal'] if boat and seg=='CAM' else NAN,
                longitud_CAM_para_tarifa_m=NAN,
                tarifa_CAM_16m_referencia_red_CLP=cfg['longitud_CAM_referencia_red_m']*cfg['tarifas_ARISTO_tabla52_p121']['metro_lineal'] if boat and seg=='CAM' else NAN,
                tarifa_pasajero_adulto_general_CLP=cfg['tarifas_ARISTO_tabla52_p121']['adulto_general'] if boat else NAN,
                numero_pasajeros_pagadores_por_vehiculo=NAN if boat and vehicle else NAN,
                costo_privado_componentes_identificados_CLP=identified,
                costo_privado_con_CAM16m_referencia_CLP=cov+cfg['longitud_CAM_referencia_red_m']*cfg['tarifas_ARISTO_tabla52_p121']['metro_lineal'] if boat and seg=='CAM' else NAN,
                costo_monetario_privado_total_completo_CLP=NAN,
                base_precios_tarifa='ARISTO_tabla52_verano2026; suma_con_COV_Dic2025_sin_homologacion_IPC' if boat else 'no_aplica',
                pasos_fronterizos=';'.join(borders),
                fuente_distancia='Costos_Red_Vial_Actual LONGITUD; cotejo GIS geodesico',
                fuente_tiempo='LONGITUD/VEL_VL_o_VP; HDM4_flujo_libre; navegacion2h_ARISTOtabla38',
                fuente_COV='Costos_Red_Vial_Actual Q/R=ROUND(tarifa_km_Tipos*LONGITUD); excluye TRANSB y conectores',
                fuente_tarifa=cfg['fuente_tarifas'] if boat else 'sin_tarifa_transbordador_por_definicion',
                fuente_anticipacion=cfg['fuente_anticipacion'] if boat else 'no_aplica',
                fuente_horarios_frontera=';'.join(cfg['fronteras'][p]['fuente'] for p in borders),
                horarios_frontera_referencia=';'.join(p+':'+cfg['fronteras'][p]['horario_ref'] for p in borders),
                condicion_tarifa_pasajero='adulto_general_sin_beneficios_no_promedio_muestral' if boat and not vehicle else 'ocupantes_pagadores_no_identificados' if boat else 'no_aplica',
                confianza_distancia='media_tramos_red; baja_para_puerta_a_puerta',
                confianza_tiempo='construido_flujo_libre; no_medido; rampas40_y94kmh_no_operacion_real',
                confianza_navegacion='documental_ARISTO_y_operador2h' if boat else 'no_aplica',
                confianza_anticipacion='regla_oficial_actual; transferencia_al_periodo_historico_condicionada' if boat else 'no_aplica',
                confianza_desembarque='no_identificado' if boat else 'no_aplica',
                confianza_H='calendario_referencia2026_y_uniformidad_horario_deseado_supuesta' if boat and s.calendario_estacion_completo else 'sin_calendario_completo; 168/f_solo_simplificacion' if boat else 'no_aplica',
                confianza_COV='privado_modelado; VP_mezcla_CS_CA; transferencia_parametros_CH_a_AR' if arg else 'privado_modelado; VP_mezcla_CS_CA_no_camion_individual',
                confianza_tarifa='tarifa_principal_documentada; ocupacion_tasas_y_longitud_CAM_pendientes' if boat else 'no_aplica',
                faltantes=';'.join(pending),estado_disponibilidad='atributos_referencia_construidos; viaje_individual_y_componentes_operativos_incompletos')


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    cfg=json.loads((ROOT/'data/templates/config_atributos.json').read_text())
    protected=[p for folder in ['matrices_od_referencia','expansion_temporal_terrestre'] for p in (ROOT/'outputs'/folder).iterdir() if p.is_file()]
    before={str(p.relative_to(ROOT)):sha(p) for p in protected}
    costpath=BASE/'Cap 7/HDM4/Costos_Red_Vial_Actual.xlsx'
    frame=pd.read_excel(costpath,sheet_name='Costos_Red_Vial')
    points=gpd.read_file(BASE/'Cap 7/SIG Red_Vial_Final/Nodos_RedVial.shp').set_index('ID_Nodo')
    assert frame.ID_Arco.is_unique
    arc_senses=pd.read_excel(BASE/'Cap 7/Arcos con sentido.xlsx')
    borderrows=arc_senses.loc[arc_senses['Unnamed: 1'].astype(str).str.contains('Chile - Argentina|Argentina - Chile')]
    pairs={tuple(map(int,v.split('-'))) for v in borderrows['Unnamed: 3'].dropna()}
    graphs={seg:(grafo_vial(frame,seg,'Chile',pairs),grafo_vial(frame,seg,'Argentina',pairs)) for seg in ['VL','CAM']}
    dep,intervals,calmetrics,monthly,seasons=horarios(cfg)
    for d,n in [(dep,'salidas_calendarios_fuente'),(intervals,'intervalos_zarpes_ciclicos'),(calmetrics,'estadisticos_calendarios'),(monthly,'horario_mensual'),(seasons,'horario_estacional')]:save(d,n)
    attrs=[];route_rows=[];candidates=[];route_objects={};route_metrics={};used=set();counter=0
    for origin,destination in cfg['OD']:
        for seg in cfg['segmentos']:
            for alt in ['terrestre_chile','barcaza_actual','terrestre_argentina']:
                border_names=[]
                if alt=='terrestre_argentina':
                    options=[]
                    for entry in cfg['fronteras_reingreso_candidatas']:
                        edges=construir_ruta(frame,graphs,cfg,origin,destination,seg,alt,entry)
                        met=medir_ruta(edges,seg,cfg)
                        permitted=seg in cfg['fronteras'][entry]['segmentos'] and seg in cfg['fronteras']['Jeinimeni']['segmentos']
                        shown=met.copy()
                        if seg=='PEATON':
                            shown['tiempo_conduccion_red_h']=NAN;shown['COV_terrestre_ARISTO_CLP']=NAN
                        candidates.append(dict(origen=origin,destino=destination,segmento=seg,pasos='Jeinimeni;'+entry if origin=='Chile Chico' else entry+';Jeinimeni',
                                               habilitado_para_segmento=permitted if seg!='PEATON' else NAN,
                                               modo_peaton_no_identificado=seg=='PEATON',**shown,seleccionado=False))
                        if permitted or seg=='PEATON' and entry=='Huemules':options.append((met['tiempo_conduccion_red_h'],entry,edges,met))
                    _,entry,edges,met=min(options,key=lambda p:p[0])
                    border_names=['Jeinimeni',entry] if origin=='Chile Chico' else [entry,'Jeinimeni']
                    for c in candidates:
                        if c['origen']==origin and c['destino']==destination and c['segmento']==seg and c['pasos']==';'.join(border_names):c['seleccionado']=True
                else:
                    edges=construir_ruta(frame,graphs,cfg,origin,destination,seg,alt);met=medir_ruta(edges,seg,cfg)
                counter+=1;rid=f'R{counter:03d}';route_objects[rid]=edges;route_metrics[rid]=met
                used.update(int(e['ID_Arco']) for e in edges)
                for ordinal,e in enumerate(edges,1):
                    boat=e['Clave']=='TRANSB'; col='VEL VP' if seg=='CAM' else 'VEL VL'
                    route_rows.append(dict(ruta_id=rid,origen=origin,destino=destination,segmento=seg,alternativa=alt,orden=ordinal,
                                           ID_Arco=int(e['ID_Arco']),nodo_desde=e['desde'],nodo_hasta=e['hasta'],sentido_respecto_NodoAB='AB' if e['desde']==e['Nodo_A'] else 'BA',
                                           ruta_fuente=e['RUTA'],region_pais_fuente=e['REGI_N'],tipo_tramo='navegacion' if boat else 'acceso_rampa' if e['RUTA']=='RAMPA' else 'carretera',
                                           longitud_fuente_km=e['LONGITUD'],velocidad_red_kmh=e[col] if seg!='PEATON' else NAN,
                                           tiempo_tramo_h=cfg['barcaza']['navegacion_h'] if boat else e['LONGITUD']/e[col] if seg!='PEATON' else NAN,
                                           COV_terrestre_aplicado_CLP=0 if boat else e['CO VP' if seg=='CAM' else 'CO VL'] if seg!='PEATON' else NAN,
                                           CO_VL_fuente=e['CO VL'],CO_VP_fuente=e['CO VP'],no_suma_CO_TRANSB_como_COV=boat,
                                           supuesto_sentido=cfg['sentidos_viales'],nivel='publico_red_NO_microdato'))
                for station in cfg['estaciones']:
                    attrs.append(fila_atributos(origin,destination,seg,station,alt,rid,met,border_names,seasons,cfg))
        # Primer producto revisable antes de extender a los otros OD prioritarios.
        if origin=='Chile Chico' and destination=='Coyhaique':
            save(pd.DataFrame(attrs),'piloto_Chile_Chico_Coyhaique')
            print('Piloto Chile Chico → Coyhaique generado; continuando sentidos inversos y Puerto Ibáñez.',flush=True)
    attrs=pd.DataFrame(attrs);routes=pd.DataFrame(route_rows);candidates=pd.DataFrame(candidates)
    attrs['secuencia_arcos']=attrs.ruta_id.map({k:';'.join(str(int(e['ID_Arco'])) for e in v) for k,v in route_objects.items()})
    attrs['secuencia_nodos']=attrs.ruta_id.map({k:';'.join(str(x) for x in [v[0]['desde']]+[e['hasta'] for e in v]) for k,v in route_objects.items()})
    save(attrs,'atributos_alternativas');save(routes,'rutas_od');save(candidates,'rutas_argentina_candidatas')
    save(attrs.loc[attrs.origen.eq('Chile Chico')&attrs.destino.eq('Coyhaique')],'piloto_Chile_Chico_Coyhaique')
    public_points=[]
    for place,p in cfg['puntos_referencia'].items():
        node=p['nodo_fisico'];public_points.append(dict(localidad_etiqueta=place,nodo=node,zona_fuente=p['zona_ARISTO'],
            x_publica=points.loc[node,'x'],y_publica=points.loc[node,'y'],descripcion=p['descripcion'],fuente=cfg['fuente_puntos']))
    save(pd.DataFrame(public_points),'puntos_referencia_OD')
    cat=gpd.read_file(BASE/'SIG/Coberturas/ARISTO/Transbordadores.shp')
    cat=cat.loc[cat.Nombre_Ser.astype(str).str.contains('Carrera',case=False)].drop(columns='geometry')
    save(cat,'catastro_transbordador_ARISTO_conservado')
    save(pd.DataFrame([dict(paso=p,segmentos_habilitados=';'.join(v['segmentos']),horario_referencia=v['horario_ref'],
                           condicion=v['habilitacion'],fuente=v['fuente'],vigencia='referencia_oficial; no_calendario_historico_estacional_completo') for p,v in cfg['fronteras'].items()]),'fronteras_habilitacion_y_horarios')
    # Conectores virtuales preservados como evidencia de localizacion pendiente, nunca10kmreales.
    connectors=[]
    for place,p in cfg['puntos_referencia'].items():
        if p['conector_ARISTO'] is not None:
            r=frame.loc[frame.ID_Arco.eq(p['conector_ARISTO'])].iloc[0]
            connectors.append(dict(OD=place,ID_Arco=int(r.ID_Arco),zona=p['zona_ARISTO'],nodo_anclaje=p['nodo_fisico'],
                                   longitud_convencional_fuente=r.LONGITUD,velocidad_fuente=r['VEL VL'],CO_fuente=r['CO VL'],
                                   distancia_real_km=NAN,tiempo_real_h=NAN,COV_real_CLP=NAN,tratamiento='excluido_suma_fisica; pendiente_localizacion_intrazonal'))
        else:connectors.append(dict(OD=place,zona=NAN,nodo_anclaje=p['nodo_fisico'],tratamiento='acceso_rampa_referencia; no_se_asigna_comuna_RioIbanez',distancia_real_km=NAN,tiempo_real_h=NAN))
    save(pd.DataFrame(connectors),'conectores_y_alcance_territorial')
    formula,rep,params=auditar_costos(frame,used,costpath)
    save(formula,'formulas_y_valores_atributos');save(rep,'reproduccion_costos_arcos')
    (OUT/'parametros_privados_fuente.json').write_text(json.dumps(params,ensure_ascii=False,indent=2))
    with ZipFile(costpath) as z:
        links=[dict(parte=n,contenido=z.read(n).decode()) for n in z.namelist() if 'externalLink' in n and n.endswith('.rels')]
    save(pd.DataFrame(links),'vinculos_externos_libro_costos')
    # Longitudes de arcos modelados contra geometria; no alteracion de la fuente.
    geo=gpd.read_file(BASE/'Cap 7/SIG Red_Vial_Final/Red_Vial_Final.shp');geod=Geod(ellps='WGS84');qc=[]
    for r in geo.loc[geo.ID_Arco.isin(used)].itertuples():
        km=geod.geometry_length(r.geometry)/1000
        qc.append(dict(ID_Arco=r.ID_Arco,longitud_fuente_km=r.LONGITUD,longitud_geodesica_km=km,razon_geometria_fuente=km/r.LONGITUD if r.LONGITUD else NAN))
    geo_check=pd.DataFrame(qc);save(geo_check,'control_distancias_GIS')
    # Validar que la red/CO se pueden reutilizar sin trasladar rutas interregionales al OD local.
    con=sqlite3.connect(BASE/'Cap 7/SIG Red_Vial_Final/SIG_RUTAS_OD/Rutas_Red_Vial_Final.gpkg')
    rows=con.execute("select RUTA,Origen,Destino,ID_Arcos,CO_VL,CO_VP from rutas where RUTA in ('1-31 Ruta 2','41-55 Ruta 1')").fetchall();con.close()
    lookup=frame.set_index('ID_Arco');macro_checks=[]
    for name,o,d,ids,co1,co2 in rows:
        sequence=[int(i) for i in ids.split(',')];node=o;reverse=0;continuous=True
        for i in sequence:
            r=lookup.loc[i]
            if node==r.Nodo_A:node=int(r.Nodo_B)
            elif node==r.Nodo_B:node=int(r.Nodo_A);reverse+=1
            else:continuous=False;break
        macro_checks.append(dict(ruta_fuente=name,continuidad=continuous,destino_correcto=node==d,
                                 arcos_recorridos_BA=reverse,CO_VL_reconstruido=lookup.loc[sequence,'CO VL'].sum(),CO_VL_fuente=co1,
                                 CO_VP_reconstruido=lookup.loc[sequence,'CO VP'].sum(),CO_VP_fuente=co2,
                                 uso='solo_comprobacion_metodo_red; NO_sustitucion_OD_local'))
    save(pd.DataFrame(macro_checks),'comprobacion_rutas_ARISTO_existentes')
    missing=[]
    for r in attrs.itertuples():
        for item in r.faltantes.split(';'):
            missing.append(dict(origen=r.origen,destino=r.destino,segmento=r.segmento,estacion=r.estacion,alternativa=r.alternativa,
                                atributo_pendiente=item,efecto='limita_atributo_completo_o_identificacion_logit; no_se_calibra',ruta_id=r.ruta_id))
    for place in ['Río Ibáñez']:
        missing.append(dict(origen='Chile Chico',destino=place,atributo_pendiente='comuna_sin_localidad; no_asignar_PuertoIbanez',efecto='fuera_reconstruccion_puntual_hasta_localizar'))
    missing.append(dict(segmento='MOTO/BICI',atributo_pendiente='velocidad_y_COV_especificos; no_se_reutiliza_VL_sin_verificar',efecto='segmentos_pendientes_de_etapa_posterior'))
    save(pd.DataFrame(missing),'faltantes_logit')
    sources=fuentes(cfg,costpath,params)
    save(sources,'fuentes_atributos')
    (OUT/'config_utilizada.json').write_text(json.dumps(cfg,ensure_ascii=False,indent=2))
    ordinary_ok=all(set(intervals.loc[intervals.calendario.eq(k),'salida_h_semana'])==
        {float(d)*24+float(h) for d,h,danger in v['salidas'] if not danger} for k,v in cfg['calendarios'].items())
    borderpairs={frozenset(p) for p in pairs}
    borders_ok=all(sum(frozenset([e['desde'],e['hasta']]) in borderpairs for e in route_objects[r.ruta_id])==
        (2 if r.alternativa=='terrestre_argentina' else 0) for r in attrs.drop_duplicates('ruta_id').itertuples())
    boats_ok=all(sum(e['Clave']=='TRANSB' for e in route_objects[r.ruta_id])==
        (1 if r.alternativa=='barcaza_actual' else 0) for r in attrs.drop_duplicates('ruta_id').itertuples())
    vehicle_bar=attrs.loc[attrs.segmento.isin(['VL','CAM'])&attrs.alternativa.eq('barcaza_actual')]
    checks=dict(rutas_continuas=True,sin_arcos_repetidos=True,
                claves_OD_unicas=not attrs.duplicated(['origen','destino','segmento','estacion','alternativa']).any(),
                cuatro_estaciones=set(attrs.estacion)==set(cfg['estaciones']),
                conectores_no_contados_como_distancia=not routes.ID_Arco.isin(frame.loc[frame.Clave.eq('C'),'ID_Arco']).any(),
                CO_arcos_reproducido=bool(rep.diferencia_CO_VL.eq(0).all() and rep.diferencia_CO_CAM.eq(0).all()),
                CG_arcos_reproducido_solo_diagnostico=bool(rep.diferencia_CG.eq(0).all()),
                CAM_Argentina_no_usa_Palavicini=not attrs.loc[attrs.segmento.eq('CAM')&attrs.alternativa.eq('terrestre_argentina'),'pasos_fronterizos'].str.contains('Palavicini').any(),
                sin_COV_vehiculo_para_peaton=bool(attrs.loc[attrs.segmento.eq('PEATON'),'COV_terrestre_privado_ARISTO_CLP'].isna().all()),
                sin_desembarque_inventado=bool(attrs.loc[attrs.alternativa.eq('barcaza_actual'),'t_desembarque_h'].isna().all()),
                anticipacion_no_es_acceso=bool((vehicle_bar.t_acceso_modelo_h-vehicle_bar.tiempo_conduccion_terrestre_h).abs().lt(1e-10).all() and vehicle_bar.t_embarque_previo_h.eq(2).all()),
                sin_suma_doble_conduccion_acceso=bool(((attrs.loc[attrs.segmento.isin(['VL','CAM']),'t_acceso_modelo_h']+attrs.loc[attrs.segmento.isin(['VL','CAM']),'t_viaje_modelo_h']-attrs.loc[attrs.segmento.isin(['VL','CAM']),'t_desplazamiento_fisico_h']).abs()<1e-10).all()),
                sin_carga_peligrosa_en_intervalos=ordinary_ok,
                numero_cruces_frontera_correcto=borders_ok,numero_transbordadores_correcto=boats_ok,
                geometria_confirma_km_en_arcos_usados=bool(geo_check.razon_geometria_fuente.between(.99,1.01).all()),
                alta11_sin_calendario_inventado=bool(monthly.loc[monthly.f_ordinario.eq(11),'H_calendario_h'].isna().all()),
                matriz_demanda_anterior_intacta=all(sha(ROOT/p)==v for p,v in before.items()))
    assert all(checks.values()),checks
    save(pd.DataFrame([dict(archivo=k,sha256=v,intacto=True) for k,v in before.items()]),'integridad_matrices_demanda')
    (OUT/'validacion.json').write_text(json.dumps(checks,indent=2))
    report(attrs,routes,candidates,seasons,calmetrics,connectors,params,checks,cfg)
    mapa(route_objects,attrs,geo,points,cfg)
    print('Atributos generados:',len(attrs),'filas;',len(route_objects),'rutas;',OUT)


def fuentes(cfg,costpath,params):
    records=[]
    sourcefiles=[('Cap 7/HDM4/Costos_Red_Vial_Actual.xlsx','Costos_Red_Vial y Tipos de arcos','km,km/h,CLP Dic2025','distancia;velocidad;COV;formula CG no usada'),
                 ('Cap 7/HDM4/COPER HDM4 PRIVADO Dic 2025.xls','referencia externa por nombre; no recalculado','precios privados Dic2025','archivo presente; fuentescache Tipos; no recalculo motorHDM4'),
                 ('Cap 7/SIG Red_Vial_Final/Red_Vial_Final.shp','geometria y atributos','EPSG4326','verificar geometria arcos seleccionados'),
                 ('Cap 7/SIG Red_Vial_Final/Nodos_RedVial.shp','ID_Nodo/x/y','grados','anclajes publicos red'),
                 ('Cap 7/SIG Red_Vial_Final/Nodos_Puntos_Pasada.xlsx','nodos_puntos_pasada','nodos','terminales y zonas'),
                 ('Cap 7/Arcos con sentido.xlsx','Hoja1','pares dirigidos','sentidos pasos y barcaza'),
                 ('Cap 10/Resumen_Rutas_Red_Vial_Actual.xlsx','Rutas Red Actual','costos/tiempos por ruta interregional','no sustituye OD local'),
                 ('Cap 7/SIG Red_Vial_Final/SIG_RUTAS_OD/Rutas_Red_Vial_Final.gpkg','rutas','ID_Arcos','validacion de suma costos y continuidad'),
                 ('SIG/Coberturas/ARISTO/Transbordadores.shp','catastro servicio Lago Gral Carrera','2h; horarios','N_VIAJES9 y omision de sabado no compatibles con8/11confirmados'),
                 ('Informe ARISTO.pdf','tablas38p112,52p121,63p131; §7p226–227; §10p297','privados,viaje completo','metodologia y atributos'),
                 ('INFORME_2_EVALSO-1.pdf','§3.2–3.3 de version local','componentes temporales y monetarios privados','separacion de tiempos y eleccion futura')]
    for file,sheet,unit,use in sourcefiles:
        p=BASE/file;records.append(dict(fuente=file,hoja_seccion=sheet,unidad=unit,uso=use,sha256=sha(p),nivel='original_ARISTO_o_EVALSO'))
    for attribute,url in [('anticipacion',cfg['fuente_anticipacion']),('calendario_y_tarifa_actual',cfg['fuente_calendarios']),('fronteras_restricciones',cfg['fronteras']['Jeinimeni']['fuente']),('Huemules_vehiculos',cfg['fronteras']['Huemules']['fuente']),('frontera_verano',cfg['fronteras']['Huemules']['fuente_verano'])]:
        records.append(dict(fuente=url,hoja_seccion=attribute,unidad='h_o_CLP_segun_atributo',uso='consulta_oficial2026-10-25; vigencia_historica_no_automatica',nivel='fuente_oficial_web'))
    scalar=[('w_horario',cfg['w_horario'],'adimensional','usuario confirma EVALSO','supuesto_exploratorio',cfg['fuente_w']),
            ('frecuencia_alta_ordinaria',11,'zarpes/semana/sentido','instruccion_usuario','confirmado_usuario',cfg['fuente_frecuencia']),
            ('frecuencia_baja_ordinaria',8,'zarpes/semana/sentido','instruccion_usuario','confirmado_usuario',cfg['fuente_frecuencia']),
            ('navegacion',2,'h','Informe ARISTO.pdf','documental','tabla38 p112; catastro2h'),
            ('anticipacion_vehiculos',2,'h',cfg['fuente_anticipacion'],'requisito_actual_transferencia_historica_condicionada','FAQ presentacion'),
            ('anticipacion_peaton',1,'h',cfg['fuente_anticipacion'],'requisito_actual_transferencia_historica_condicionada','FAQ presentacion'),
            ('desembarque',NAN,'h','no_identificado','faltante','no_anadir_otra_operacion_por_defecto'),
            ('tarifa_VL_hasta5_5m',21270,'CLP/vehiculo/tramo','Informe ARISTO.pdf','documentado_historico','tabla52 p121'),
            ('tarifa_adulto_historica',2550,'CLP/persona/tramo','Informe ARISTO.pdf','documentado_historico','tabla52 p121'),
            ('tarifa_CAM_metro',7310,'CLP/m/tramo','Informe ARISTO.pdf','documentado_historico','tabla52 p121'),
            ('tarifa_adulto_actual_contraste',2510,'CLP/persona/tramo',cfg['fuente_calendarios'],'referencia_actual_no_sustituye_historica','web25/10/2026'),
            ('CAM_longitud_referencia_red',cfg['longitud_CAM_referencia_red_m'],'m','Costos_Red_Vial_Actual.xlsx','supuesto_implicito_red_NO_promedio_muestral',cfg['fuente_longitud_CAM_referencia']),
            ('peso_CS_COV_VP',params['peso_CS'],'proporcion','Costos_Red_Vial_Actual.xlsx','parametro_fuente','Tipos de arcos W3'),
            ('peso_CA_COV_VP',params['peso_CA'],'proporcion','Costos_Red_Vial_Actual.xlsx','parametro_fuente','Tipos de arcos V3'),
            ('VST_CG_ARISTO_no_usado_para_logit',params['VST_CG_fuente_no_usado'],'CLP/h','Costos_Red_Vial_Actual.xlsx','solo_auditoria_no_regresor','Tipos V15'),
            ('fecha_precios_privados',params['fecha_precios'],'fecha','Costos_Red_Vial_Actual.xlsx','documentada','Tipos V44; U45/V45 precio privado'),
            ('longitud_conector_convencional',10,'km_convencionales_NO_fisicos','Costos_Red_Vial_Actual.xlsx','excluida_distancia_y_tiempo_reales','Conect39 yConect31')]
    for name,value,unit,source,status,section in scalar:
        records.append(dict(parametro=name,valor=value,unidad=unit,fuente=source,estado=status,hoja_seccion=section,nivel='parametro_identificado'))
    return pd.DataFrame(records)


def report(attrs,routes,candidates,seasons,calmetrics,connectors,params,checks,cfg):
    pilot=attrs.loc[attrs.origen.eq('Chile Chico')&attrs.destino.eq('Coyhaique')&attrs.estacion.eq('Verano')]
    metriccols=['segmento','alternativa','distancia_terrestre_total_km','tiempo_conduccion_terrestre_h','tiempo_navegacion_h',
                'anticipacion_exigida_h','COV_terrestre_privado_ARISTO_CLP','tarifa_transbordador_principal_CLP',
                'costo_privado_componentes_identificados_CLP']
    temporal=attrs.loc[attrs.origen.eq('Chile Chico')&attrs.destino.eq('Coyhaique')&attrs.segmento.eq('VL')&attrs.estacion.eq('Invierno')]
    seq=[]
    for r in pilot.loc[pilot.segmento.isin(['VL','CAM'])].itertuples():
        arcs=routes.loc[routes.ruta_id.eq(r.ruta_id)].sort_values('orden');labels=arcs.ruta_fuente.tolist()
        compact=[v for i,v in enumerate(labels) if i==0 or v!=labels[i-1]]
        seq.append(dict(ruta_id=r.ruta_id,segmento=r.segmento,alternativa=r.alternativa,n_arcos=len(arcs),
                        carreteras_en_orden=' → '.join(compact),pasos=r.pasos_fronterizos or 'sin_frontera'))
    others=attrs.loc[~(attrs.origen.eq('Chile Chico')&attrs.destino.eq('Coyhaique'))&attrs.segmento.isin(['VL','CAM'])&attrs.estacion.eq('Verano')]
    text=f'''# Atributos de alternativas de referencia — piloto Chile Chico → Coyhaique

## Alcance y lectura de los resultados

Se reconstruyen cuatro OD prioritarios y tres alternativas, por VL, CAM y peatones, manteniendo las cuatro estaciones. Las rutas llegan de origen a destino entre **puntos de referencia públicos explícitos de la red**: Chile Chico1358, Coyhaique1223 y acceso de Puerto Ibáñez1021. No son domicilios individuales ni centros urbanos validados. Para Chile Chico/Coyhaique son los anclajes físicos de las zonas comunales39/31. Puerto Ibáñez se ancla por el acceso a su rampa; **no se renombra la comuna Río Ibáñez ni se usa su centroide40**.

La matriz contiene el recorrido completo de red entre esos anclajes y sus limitaciones de localización intrazonal. Los conectores ARISTO39/31 tienen10km convencionales, velocidad0 y CO0: no se convierten en10km físicos ni en tiempos reales nulos. El viaje puerta a puerta individual sigue sin identificar. Todos los tramos terrestres, incluyendo el acercamiento y salida de terminales, están en las sumas. No se trasladó una ruta interregional como sustituto de estos OD.

## Tabla detallada Chile Chico → Coyhaique

Distancias en km; tiempos en h; costos COV en CLP privados diciembre2025; tarifa principal histórica de tabla52 de ARISTO (verano2026). La columna de costos identificados suma componentes documentados, **no el total monetario completo**: faltan localización intrazonal, pagos adicionales, ocupación/condiciones de pasajero y posibles tasas. CAM de barcaza requiere longitud, y peatones requieren modo/tarifa de acceso terrestre. No se asigna a peatones el COV ni el tiempo de conducción de un VL.

{md(pilot[metriccols])}

Secuencia compacta de carreteras en el mismo orden de los arcos. La secuencia completa de nodos/arcos está en cada fila de atributos y normalizada por tramo en `rutas_od.csv`:

{md(pd.DataFrame(seq))}

Desglose temporal del piloto VL en invierno como ejemplo del calendario de baja de referencia; las horas de conducción ya están dentro del acceso de la alternativa barcaza y no vuelven a sumarse:

{md(temporal[['alternativa','tiempo_conduccion_terrestre_h','t_acceso_modelo_h','t_viaje_modelo_h','anticipacion_exigida_h','t_desembarque_h','H_h','H_168_f_h','t_horario_h','t_horario_168_f_h','t_frontera_tramites_h']])}

Argentina VL selecciona el corredor Jeinimeni–Palavicini; CAM usa Jeinimeni–Huemules porque la referencia oficial restringe Palavicini a vehículos menores. No se fuerza al camión por el recorrido del VL. Ambas opciones se calcularon y se conservan en el archivo de candidatos, incluida la alternativa físicamente trazable pero no habilitada para CAM por Palavicini:

{md(candidates.loc[candidates.origen.eq('Chile Chico')&candidates.destino.eq('Coyhaique'),['segmento','pasos','habilitado_para_segmento','distancia_terrestre_total_km','tiempo_conduccion_red_h','COV_terrestre_ARISTO_CLP','seleccionado']])}

## Sentido inverso y Puerto Ibáñez: atributos calculados

Los mismos precios/velocidades de fuente se mantienen por estación como supuesto estático porque no existen parámetros viales estacionales separados. Esta tabla muestra la referencia de verano; el detalle de cuatro estaciones, horarios y pendientes está en el CSV. El sentido inverso se reconstruye y coteja por su secuencia propia.

{md(others[['origen','destino']+metriccols])}

Las rutas por Argentina quedan condicionadas a apertura, documentos, mercancía y restricciones. El comunicado oficial de referencia07/04/2026 indica08–20 para estos pasos; una comunicación de14/01/2026 extiende Huemules a08–21 entre diciembre y marzo. No se asume ese horario para todas las estaciones ni se inventan minutos de trámite. No se han sincronizado husos horarios o fechas individuales. El costo HDM4 de arcos argentinos reutiliza parámetros privados CLP de la red: no demuestra precios argentinos efectivos ni incorpora seguro/permisos internacionales.

## Partición de tiempos, sin duplicaciones

La conducción total es un atributo descriptivo calculado como suma longitud/velocidad de **todos** los arcos terrestres. En barcaza, `t_acceso_modelo_h` contiene esos acercamientos y salidas entre OD y terminal; `t_viaje_modelo_h` contiene navegación. En tierra, acceso0 y viaje=conducción. Por tanto:

`t_acceso_modelo + t_viaje_modelo = conducción_total + navegación`

La conducción descriptiva no vuelve a sumarse junto al acceso. `tiempo_acceso_rampas_modelo_h` es solamente el subconjunto de arcos RAMPA, conservado para inspección. Sus velocidades de fuente son40 y94km/h, velocidades genéricas de tipología que no certifican operación de terminal; quedan señaladas para mejorar. Los accesos de esta red llegan al terminal, no representan maniobras sobre cubierta.

La [FAQ oficial de Naviera Austral]({cfg['fuente_anticipacion']}) exige1h de presentación para pasajeros y2h para vehículos/rodados. Se utiliza como referencia empírica de anticipación, con vigencia histórica no demostrada. No es acceso terrestre ni espera aleatoria medida. Las operaciones previas se consideran contenidas en esa anticipación; no se agrega otra media hora de embarque. **Desembarque no identificado**: `t_embarque_desembarque_completo_h` y total físico completo de la barcaza permanecen faltantes. La navegación principal es2h (ARISTOtabla38 y operador); la división37,277km/18kmh de la red da otro valor y se conserva solo como diagnóstico, no se suma una segunda navegación.

## Horario y cuatro estaciones

Se recuperaron las salidas ordinarias, se excluyeron las marcadas como carga peligrosa y se calcularon todos los intervalos, incluido el cierre domingo→lunes. La media aritmética de intervalos es por identidad168/f; no añade información de distribución. Para el contraste de calendario se define `H=sum(Δh²)/168`, que pondera cada intervalo por su duración. Supone horario deseado uniforme durante la semana y **no es espera en muelle**. La penalización exploratoria es0,25H; se entrega también0,25×168/f.

{md(calmetrics)}

El calendario de baja publicado desde01/05/2026 contiene8ordinarios por sentido y permite calcular H de referencia distinto por sentido. La transferencia a otras campañas, en particular primavera2025, no se presenta como calendario histórico verificado. El de alta publicado contiene12ordinarios por sentido, y ARISTOtabla38 también presenta12; el usuario confirmó11. No se elimina una salida arbitrariamente: para11solo está identificado el escenario simplificado168/11. La página dice literalmente31deabril2026, fecha inválida que no se usa como vigencia comprobada. El catastro SIG de ARISTO informa9semanales y su texto omite un sábado desde Chile Chico; se conserva como discrepancia, no se emplea para inventar el calendario11.

{md(seasons)}

Otoño se mantiene desagregado en marzo/abril11 y mayo8. Los atributos estacionales simplificados son promedios ponderados por días de cada mes de2026; no se calcula168 dividido por una frecuencia media que mezcle regímenes. Cuando falta parte del calendario de una estación, el H de calendario estacional queda NaN y se informa la cobertura parcial de días. Las rutas terrestres tienen penalización de servicio0; los trámites y restricciones horarias de frontera siguen como componentes separados faltantes.

## Fórmulas, precios privados y tarifas

El libro Costos_Red_Vial verifica `CO_VL=ROUND(costo_km_VL×LONGITUD,0)` y la fórmula equivalente CAM. `Tipos de arcos!V44` documenta diciembre2025 y U45/V45 precios privados; el enlace externo apunta a `COPER HDM4 PRIVADO Dic 2025.xls`, presente en los anexos. Se reproducen los cálculos de los arcos utilizados con sus valores guardados, sin recalcular Excel ni el motor HDM4 externo. VP es una mezcla de CS/CA con pesos de fuente {params['peso_CS']:.6f}/{params['peso_CA']:.6f}; no es el costo exacto de cada camión ni se sustituye por VL.

`CostGen VL` contiene CO más valoración temporal de fuente, con VST={params['VST_CG_fuente_no_usado']:.0f} y multiplicador1,5 para no pavimentado en las fórmulas terrestres. **No se usa como desembolso ni como tiempo físico**, ni se suma a CO nuevamente. Tampoco se utiliza ese VST como coeficiente calibrado del proyecto. El ruteo local minimiza tiempo libre con velocidades propias VL/VP.

El CO del arco de transbordador es una tarifa/paquete, no operación vehicular terrestre. En la red VL={params['tarifa_VL_embebida_red']:.0f} coincide con21270+2×2550; es una composición implícita de dos pasajes que no se impone a todo vehículo. CAM={params['tarifa_CAM_embebida_red']:.0f} coincide con7310×16m; se conserva el ejemplo16m como referencia, sin fijarlo como longitud de todos los camiones. Las tarifas principales usadas se toman por separado de ARISTOtabla52: VL21270 hasta5,5m, adulto2550, CAM7310por metro. La web actual publica2510adulto, discrepancia que no sustituye silenciosamente el antecedente histórico. No se deduce número de pasajeros pagadores ni se cobra otra vez al conductor desde una regla no verificada para el período. El tarifario web advierte posibles tasas de embarque no incluidas; las tasas faltan.

Los costos de componentes identificados son útiles para revisar magnitudes, pero la suma de COVdiciembre2025 y tarifasverano2026 no está homologada por IPC a una base común; esa homologación queda pendiente. No se convierten precios sociales a privados ni se utiliza la valoración social de EVALSO para la elección.

## Trazabilidad, verificaciones y pendientes

`rutas_od.csv` contiene cada ID de arco, orden, nodo inicial/final, orientación AB/BA, ruta, país/región, distancia, velocidad, tiempo y COV aplicado. `conectores_y_alcance_territorial.csv` conserva los conectores omitidos de la suma física. Se calculó cada sentido independientemente. Cuando los costos terrestres coinciden, es porque la red de fuente no tiene gradientes/sentidos asimétricos; no significa igualdad observada. El archivo de arcos con sentido corrobora ambos pares para pasos y barcaza; la legalidad de sentidos únicos urbanos y restricciones de peso está pendiente.

Se cotejaron longitudes con geometría geodésica y se reprodujeron dos rutas interregionales solo como comprobación del método de red y suma de costos; no entran como OD locales. Los resultados de demanda de pasos07/08 se verifican por hash y no se modifican. Validaciones: `{json.dumps(checks,ensure_ascii=False)}`.

Para calibración faltan: accesos intrazonales y definición del OD individual, atributos de modos de acceso de peatones, desembarque, calendario11, preferencias horarias, tarifas efectivamente pagadas y longitud/ocupantes por vehículo, trámite y costos de frontera, restricciones por camión, parámetros privados argentinos y variación estacional de velocidades/cierres. Río Ibáñez comunal permanece sin localización puntual. Motos/bicicletas requieren velocidades y COV propios; esta primera matriz se limita a VL/CAM/PEATON. Los demás OD de las matrices exploratorias no se amplían antes de revisar este piloto.

Archivos requeridos: `atributos_alternativas.csv`, `rutas_od.csv`, `fuentes_atributos.csv`, `faltantes_logit.csv`; primer detalle `piloto_Chile_Chico_Coyhaique.csv`. Parámetros: `data/templates/config_atributos.json`. Reproducción: `python scripts/09_construir_atributos.py`. Sin calibración, probabilidades, expansión de demanda, SBO, commit ni push.
'''
    (OUT/'resumen_atributos.md').write_text(text,encoding='utf-8')


def mapa(route_objects,attrs,geo,points,cfg):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    chosen=attrs.loc[attrs.origen.eq('Chile Chico')&attrs.destino.eq('Coyhaique')&attrs.segmento.eq('VL')&attrs.estacion.eq('Verano')]
    fig,ax=plt.subplots(figsize=(10,8));colors={'terrestre_chile':'#d68b00','barcaza_actual':'#1879b5','terrestre_argentina':'#9256aa'}
    for r in chosen.itertuples():
        ids=[int(e['ID_Arco']) for e in route_objects[r.ruta_id]]
        g=geo.loc[geo.ID_Arco.isin(ids)].to_crs(32718)
        g.plot(ax=ax,color=colors[r.alternativa],linewidth=2.2,label=r.alternativa)
    for name,node in [('Chile Chico',1358),('Coyhaique',1223),('Puerto Ibáñez\n(acceso rampa)',1021),('Jeinimeni',1323),('Palavicini',1311)]:
        p=gpd.GeoSeries([points.loc[node,'geometry']],crs=4326).to_crs(32718).iloc[0]
        ax.plot(p.x,p.y,'o',color='#202b32',ms=4);ax.annotate(name,(p.x,p.y),xytext=(5,5),textcoords='offset points',fontsize=8)
    ax.set_aspect('equal');ax.set_xlabel('Este UTM18S (m)');ax.set_ylabel('Norte UTM18S (m)')
    ax.set_title('Chile Chico → Coyhaique: rutas de referencia VL\nArcos ARISTO; puntos representativos, sin conexiones domiciliarias')
    ax.legend(loc='lower left');ax.grid(alpha=.15);fig.tight_layout();fig.savefig(OUT/'mapa_rutas_piloto.png',dpi=180);plt.close(fig)


if __name__=='__main__':main()
