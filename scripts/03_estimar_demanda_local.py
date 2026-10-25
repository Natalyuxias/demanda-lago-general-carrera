"""Estimación condicional local. No escribe fuentes ni exporta microdatos.

Uso: python scripts/03_estimar_demanda_local.py [--anexos /ruta] [--refresh]
Las constantes y su procedencia están en data/templates/config_estimacion.json.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import math
import sys
import struct
import unicodedata
from pathlib import Path
from urllib.parse import unquote
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/estimacion_preliminar'
CACHE = ROOT / 'data/interim/estimacion_preliminar'

def norm(v):
    if pd.isna(v): return ''
    return ' '.join(''.join(c for c in unicodedata.normalize('NFKD', str(v)) if not unicodedata.combining(c)).lower().split())

def es_lago(v):
    return 'ibanez' in norm(v) and 'chile chico' in norm(v)

def sentido(v):
    s = norm(v)
    if 'chile chico' not in s or 'ibanez' not in s:
        raise ValueError(f'Sentido no reconocido: {v}')
    return 'Chile Chico a Ibanez' if s.index('chile chico') < s.index('ibanez') else 'Ibanez a Chile Chico'

def numero(v):
    return isinstance(v, (float,int)) and not isinstance(v,bool) and math.isfinite(v)

def repartir(q, n, total):
    """Sin muestra no hay participacion identificada: conservar NaN."""
    if total <= 0: return float('nan')
    if n < 0 or n > total or q < 0: raise ValueError('Insumos invalidos')
    return q*n/total

def anualizar(flujos, dias, razon_invierno):
    assert sum(dias.values()) == 365
    if any(pd.isna(flujos.get(s)) for s in ['Verano','Primavera','Normal']): return float('nan')
    return sum(flujos[s]*dias[s] for s in ['Verano','Primavera','Normal']) + flujos['Normal']*razon_invierno*dias['Invierno']

def dbf_carrera(path):
    """Solo atributos del servicio, sin dependencia SIG ni geometrías."""
    with path.open('rb') as fh:
        h=fh.read(32);n=struct.unpack('<I',h[4:8])[0];hlen,rlen=struct.unpack('<HH',h[8:12]);fields=[]
        while True:
            x=fh.read(32)
            if x[0]==13: break
            fields.append((x[:11].split(b'\0')[0].decode(),x[16]))
        fh.seek(hlen)
        for _ in range(n):
            row=fh.read(rlen);off=1;record={}
            for name,length in fields:
                record[name]=row[off:off+length].decode('utf-8',errors='replace').strip();off+=length
            if 'Carrera' in str(record):return record
    raise ValueError('Servicio General Carrera no encontrado')

def save(df, nombre):
    df.to_csv(OUT / f'{nombre}.csv', index=False, encoding='utf-8-sig', na_rep='NaN')

def mdtable(df):
    def fmt(x):
        if pd.isna(x):return 'No identificado'
        if isinstance(x,(float,int)) and not isinstance(x,bool):return f'{x:,.2f}'
        return str(x).replace('|','/').replace('\n',' ')
    return '\n'.join(['| '+' | '.join(df.columns)+' |','| '+' | '.join(['---']*len(df.columns))+' |']+['| '+' | '.join(fmt(x) for x in row)+' |' for row in df.itertuples(index=False,name=None)])

def reporte(config, comp, both, annual_both, controls, samp):
    ref=comp.groupby(['variante','temporada','segmento']).flujo_dia.agg(lambda s:s.sum(min_count=2)).reset_index()
    med=both.loc[both.escenario.eq('medio') & both.segmento.eq('VL')]
    ann=annual_both.loc[annual_both.escenario.eq('medio') & annual_both.segmento.eq('VL')]
    small=samp.groupby(['temporada','sentido','segmento']).n_od.sum().reset_index(name='encuestas')
    roads=pd.read_csv(OUT/'carretera_Cerro_Castillo_NO_demanda_local.csv')
    text=f'''# Primera estimación cuantitativa: Bahía Jara–Levicán

Proyecto Conectividad Lacustre Lago General Carrera: Bahía Jara–Mallín Grande. Entrega local del 24/10/2026. **Escenarios parciales de desvío desde el cruce histórico Puerto Ibáñez–Chile Chico; no demanda total del proyecto frente a la SBO.**

## Resultado comunicable

La cadena de cálculo original ARISTO es reproducible. Hay fundamento documental para una sensibilidad que usa media por zarpe: la sección 2.2.6 la prescribe y la tabla 78 documenta dos salidas por sentido en primavera. No se sustituye el original. Se presentan cuatro variantes de flujo base y tres tasas de captación supuestas (10/25/50%). Los resultados no son intervalos estadísticos ni probabilidades de un logit.

Esta tabla muestra flujos diarios equivalentes del cruce completo, **ambos sentidos**, separados por unidad física. VL es la clase de privados de ARISTO, que incluye motos, casas rodantes y bicicletas; CAM incluye furgones de carga y remolques según fórmulas. PEATON es columna G, sin bicicletas. ARISTO_original no proporciona expansión de peatones en TD Transb.

{mdtable(ref)}

Captación **media supuesta del 25%**, por día equivalente, ambos sentidos, clase VL. Los cuatro OD iniciales son Chile Chico↔Coyhaique y Chile Chico↔Río Ibáñez. El cruce completo es una envolvente de contraste: no se ha verificado que todos sus OD puedan utilizar Jara–Levicán.

{mdtable(med[['variante','temporada','alcance','referencia_dia','captada_barcaza_dia']])}

Los escenarios bajo y alto son respectivamente 0,4 y 2 veces el escenario medio, sin cambiar su universo ni sus supuestos. Las tablas C contienen las tres tasas, por OD, sentido, campaña y segmento. Se identificaron otros OD en el cruce y se entregan como candidatos; su factibilidad por accesos, destino final y recorrido completo no está confirmada.

## Fuentes primarias y jerarquía

* `Informe ARISTO.pdf` (nombre exacto, ubicado en la raíz de Anexos): sección 2.2.6 p.7; tabla38 pp.111–112; capítulo6, tabla78 p.173; tablas107/138/169 pp.187/207/223; capítulo8.1 pp.231–233 y 8.2 desde p.233. Se consultó el PDF local, sin fuentes web.
* `Cap 6/BBDD EOD/Base de datos agregada.xlsx`: BBDD-1, BBDD-2 (solo válidas), TD Transb. Entrevistas y cálculo de flujo marítimo, no unidades intercambiables.
* `Cap 6/Mediciones de flujo/Consolidado Flujos.xlsx`: Transbordadores y TD; Carretera se revisa para cobertura terrestre, sin asignar sus volúmenes completos a OD locales.
* `Cap 8/EOD_Puntos_Pasada_Final.xlsx`: arcos_ptocontrol, seis estratos del lago; `Cap 8/Matrices Situación Actual.xlsx`: seis hojas; `PIJA Final.xlsx`: estructura de los PIJA. Se conservan fórmulas y caches seleccionados, sin recalcular Excel.
* `SIG/Coberturas/ARISTO/Transbordadores.dbf`: atributos del servicio General Carrera, itinerarios, nueve viajes semanales, referencia sin vigencia operacional verificada. No se usó geometría personal.
* `Estudios de consultoras/Informe 2c - MOP aysen vf.pdf` (CITRA), §6.1 pp.75–90: campaña estival 20/01–05/02/2026; matrices de entrevistas sin expansión y conteos 08:00–20:00. Contraste independiente, sin sumar muestras ni convertir conteos de corredor en demanda exclusivamente terrestre.
* `INFORME_2_EVALSO-1.pdf`: versión local con demanda en capítulo3, no capítulo5 (este último trata tecnologías/capacidad). Se implementa parcialmente la secuencia de demanda y ecuación3.1; capítulo2 define SBO de dos naves y mejora de frecuencia. Cuadro8.4 aporta 10/25/50% como contraste. No se encontró otro PDF EVALSO en Estudios de consultoras.

Auditorías y CSV de variables hechos por IA se usaron solo como índices de ubicación. No son fuentes de ARISTO, evidencia de representatividad ni respaldo de supuestos. Esta entrega verifica los cálculos originales y reemplaza la conclusión anterior de detenerse sin producir escenarios numéricos.

## Fórmulas y unidades verificadas

Sea C la suma del conteo por campaña/sentido/clase; Z el número de zarpes representados y f los zarpes por semana.

1. **TD Transb original**: Q_ARISTO = C × f / 7, f=9 en los seis estratos. Las celdas G20:H25 y sus caches están en `formulas_y_cache.csv`; en Python se reproducen con los conteos. La hoja TD del Consolidado contiene sumas, no medias por zarpe. El resultado es diario equivalente según el libro; C necesita interpretarse como flujo por servicio para que la unidad sea consistente con §2.2.6.
2. **Alternativas**: Q = (C/Z) × f / 7. Primavera Z=2 por sentido, respaldado por tabla78 y las dos filas. Verano Z=2 y Normal Z=1 son supuestos una-fila-un-zarpe; no hay horas verificadas para probarlos. Variantes media_zarpe_9 y media_zarpe_12 usan f=9 y f=12 respectivamente. La segunda no afirma que el itinerario de tabla38 rigiera en todas las temporadas. Un día equivalente es promedio semanal/7, no flujo de cada día operativo.
3. **Sensibilidad de versión**: media_zarpe_12_tabla169 conserva las demás campañas y usa en Normal los valores VL/CAM del PDF (16/4 Chile Chico→Ibáñez y 16/3 en sentido inverso). El Excel usa 17/5 y 16/4. Para peatones conserva columnaG del Excel para no tratar bicicletas como personas.
4. **arcos_ptocontrol**: H/I buscan C en `[3]TD` del Consolidado y multiplican por 9/7; J=F/H y K=G/I. El numerador es número acumulado de encuestas de su tabla agregada y el denominador es flujo diario calculado: la unidad del cociente es días de muestra equivalentes, no una probabilidad de inclusión automáticamente válida. Que J>1 no basta para invalidar la aritmética, pero impide interpretarlo sin exposición como fracción muestral.
5. **Mensual**: L=H×30,5; M=I×30,5; N=F/L y O=G/M. El 30,5 es un fijo de Excel, no calendario de esta anualización. Se reprodujeron los {controls['factores_reproducidos']} cálculos de flujo mensual y factores en seis estratos/clases, incluyendo vacíos cuando el numerador es cero. No se invirtieron estos factores para expandir la EOD local completa: la clasificación original de encuestas no coincide completamente con las unidades locales.
6. **Matrices**: D=E/SUM(F:BI), donde F:BI son PIJA × factor buscado. El archivo usa en fila1 los factores **mensuales** N/O de arcos_ptocontrol (columnas14/15); por dimensión, D resulta vehículos/mes, aunque §8.2 describe matrices diarias. Se reprodujeron {controls['formulas_matriz_reproducidas']} de {controls['formulas_matriz_total']} celdas D usando los denominadores almacenados, en las seis hojas de Matrices Situación Actual y las seis copias internas de EOD_Puntos_Pasada_Final. Eso verifica la aritmética del cache, no la vigencia de sus vínculos ni todos los PIJA. No se transformaron silenciosamente esas matrices ni se usaron para la matriz local. No contienen los cuatro OD iniciales (zonas39–31 y39–40, ambos sentidos); las hojas internas adicionales se revisan en el registro de cobertura.

Los vínculos apuntan a ubicaciones originales Windows; se registran destinos y candidatos locales. Las búsquedas XLOOKUP/IFERROR pueden ocultar ausencias con vacío y las matrices con cero. Openpyxl **no** recalcula fórmulas: se preservó cache, se reconstruyó en Python lo indicado y no se actualizó ningún vínculo. No se puede certificar frescura de cache solo porque la identidad aritmética cierre.

## Frecuencia y representatividad

Tabla38: lunes2, martes1, miércoles2, jueves1, viernes2, sábado2, domingo2 =12 por sentido/semana. Tabla78: 08/11/2025, Chile Chico→Ibáñez 09:00 y16:00; inverso12:00 y20:00. La tabla38 indica retorno del sábado12:00 y19:00: hay incluso diferencia horaria entre ambas tablas. Se confirma que hay dos salidas medidas, no equivalencia de itinerarios completos.

El SIG declara nueve semanales, pero su texto distingue carga peligrosa: martes16:00 desde Chile Chico y martes11:00 desde Ibáñez. Su texto enumera siete salidas ordinarias y una especial desde Chile Chico, y ocho ordinarias más una especial desde Ibáñez; no concilia plenamente con el atributo nueve para ambos sentidos. La vigencia y cobertura estacional no constan. No se asignan automáticamente servicios de carga peligrosa a peatones o vehículos comunes. Las frecuencias9/12 en tablas son sensibilidades de expansión, no bitácoras ni confirmación de viajes accesibles a todas las clases.

Primavera: dos conteos por sentido coinciden en número con los dos zarpes documentados. No se dispone de bitácora para garantizar todos los efectivamente operados. Verano: dos filas por sentido; las dos de ida CC→Ibáñez son del01/02, las inversas son27/01 y01/02. Normal: una fila por sentido el10/05. No se demuestra que cubran todos los zarpes operados. Cada alternativa conserva esta incertidumbre.

## Distribución local y segmentos

Se enlazaron ID2 únicos entre BBDD-1 y BBDD-2. Hay {controls['originales']:,} originales, {controls['validas']:,} válidas y {controls['interceptadas_cruce']} válidas interceptadas en el cruce. Los {controls['pares_iniciales']} registros de los cuatro OD iniciales incluyen entrevistas externas: para distribuir el flujo lacustre se usan **solo** las interceptadas en el cruce. El resto conserva evidencia declarada o ruta desconocida; nunca se le asignó tierra por estar en carretera.

Por estrato temporada×sentido×segmento: p_OD = n_OD / n_estrato; Q_OD = Q_estrato × p_OD. Es **supuesto provisional de composición representativa**, no ponderación validada por diseño muestral. No se mezclan peatones con vehículos en el denominador. En verano Ibáñez→Chile Chico, arcos_ptocontrol cuenta 39 encuestas VL; la selección por tipo exacto conserva 35. Las otras cuatro son respuestas de pasajeros BUS turismo (1), TRANSFER (1) y VAN turismo (2), que llevan bandera Encuesta VL=1 en BBDD-1. Se excluyen justificadamente de la composición vehicular para no representar pasajeros como vehículos; las demás celdas VL/CAM concuerdan. El enlace y diferencia están en `control_enlace_TD_arcos_y_muestra.csv`. Se conserva el flujo sin OD cuando falta muestra. Si n_estrato=0 la distribución es NaN, nunca cero. Si existe muestra pero un OD no aparece, su cero proporcional es muestral y no prueba ausencia poblacional. Los dos registros válidos con bandera Duplicada requieren protocolo; no se borraron automáticamente.

{mdtable(small)}

VL agrupa auto/camioneta/vehículo privado/moto; CAM, camión o furgón de carga; PEATON, pasajero peatón. Pasajeros de bus, van y transfer quedan fuera del flujo vehicular: una respuesta de pasajero no representa un bus/van. La clase VL de conteos incluye bicicletas (W=SUM(H:K,T:V)) mientras las etiquetas de la muestra no permiten una distribución específica de ellas: la asignación de esa clase a OD vehiculares es una aproximación declarada. Hay una bicicleta observada en primavera Ibáñez→Chile Chico y una en Normal Chile Chico→Ibáñez. Los conteos de remolques tampoco equivalen necesariamente a camiones tractores; CAM conserva la definición ARISTO Y=SUM(L,P:S). No se convierten camiones a toneladas ni peatones a todos los pasajeros.

La columnaZ de peatones suma G+V (bicicletas); el PDF declara esta inclusión. Se usa G para escenarios de personas. No se suma ocupación de autos ni grupos a peatones: no se ha validado si representan personas adicionales o unidades de respuesta. No se informa carga en toneladas.

No hay muestra válida de camiones CC→Ibáñez en verano aunque hay conteo positivo: su flujo queda sin asignar por OD. Tampoco hay peatones válidos Ibáñez→CC en verano. Por ello algunos resultados locales y anuales son NaN. La envolvente del cruce sigue siendo cuantificable por conteos bajo sus supuestos, sin asignar OD inventados.

## Anualización condicional, no proyección temporal

Se usa año sintético no bisiesto, sin crecimiento: verano90 días, primavera91, Normal92 e invierno92, sin solapamiento. Son supuestos del ejercicio; Normal no se renombra invierno. Para invierno se sensibiliza 50% o100% del flujo Normal, manteniendo por supuesto la misma composición OD. A_anual =90Q_verano+91Q_primavera+92Q_Normal+92×r_inv×Q_Normal. Si falta un estrato, el resultado se mantiene NaN. No se multiplicó una campaña por365 ni se atribuyó crecimiento a la segunda nave.

Captación anual al25%, ambos sentidos, clase VL (unidades/año):

{mdtable(ann[['variante','alcance','invierno','captada_barcaza_ano']])}

## SBO, tierra, inducción y pendientes

EVALSO define dos naves, optimización de itinerarios y coordinación tarifaria, sin parrilla completa verificable para esta aplicación. No se duplica la demanda por duplicar la flota. El cálculo conserva volumen histórico y aplica captaciones exógenas: es un contraste inicial **sin modelar respuesta a la SBO**. No mide todavía el cambio de cuota que produciría la mayor frecuencia de la barcaza existente. Debe presentarse así al equipo.

Demanda exclusivamente terrestre, inducida y total del proyecto quedan **NaN/no identificadas**, no cero. ARISTO ofrece conteos de corredores y entrevistas externas, pero no basta asignar todo Cerro Castillo ni todos los pasos fronterizos a los OD locales; hace falta ruta verificada y exposición compatible. CITRA publica matrices muestrales de verano y conteos12h, sin identificar individualmente quién usó lago en esos OD; no se suman a ARISTO. No se estimó una demanda terrestre expandida con esos agregados.

En ODC-3 Cerro Castillo, los conteos originales por campaña/sentido son los siguientes; describen la exposición medida del corredor, no un flujo diario local ni una elección exclusivamente terrestre:

{mdtable(roads[['Temporada','SENTIDO','intervalos','dias_etiquetados','privados_observados','camiones_observados']])}

Se conserva además TD Carr del punto. En primavera/verano divide los acumulados de 80 intervalos por dos jornadas; Normal tiene 40 intervalos y una jornada. Es promedio de la jornada observada (40 cuartos de hora), no una expansión automática a24h. La EOD local fuera del cruce se entrega en `candidatos_fuera_barcaza_NO_flujo_terrestre.csv`, manteniendo ruta desconocida cuando corresponde.

Se revisó la ecuación de costos generalizados de EVALSO (3.2) y la inducción, pero faltan atributos comparables de viaje completo histórico/SBO/proyecto, acceso Jara/Levicán, tarifa, navegación, embarque y desfase horario. No existe una reducción de CG defendible para aplicar una elasticidad de manera homogénea. No se calculan CG, inducción, preferencias ni logit con números inventados. Falta justificar transferencia de elasticidades, depurar unidades y definir año monetario. Las tasas de crecimiento interregionales tampoco se transfieren automáticamente a estos OD locales.

Prioridades de mejora: manifiestos y zarpes efectivos por campaña/sentido/clase; calendario ordinario versus carga peligrosa y cancelaciones; resolver discrepancia Normal PDF/Excel; definición de vehículo/remolque/peatón/bicicleta; protocolo EOD y no respuesta, ceros muestrales y duplicaciones; muestra local de camiones/peatones faltante; invierno y calendario representativo; itinerarios y tarifas SBO; accesos y atributos del proyecto; rutas terrestres verificadas y capacidad/rechazos.

## Archivos, reproducción y confidencialidad

Tablas A/B/C/D se encuentran en esta misma carpeta. `A_demanda_referencia_od.csv` conserva muestra, participación, fuente, variante y unidad; `B_comparacion_expansion.csv` conserva el original y diferencias; `C_captacion_por_od.csv`, `C_resumen_escenarios.csv` y `C_resumen_ambos_sentidos.csv` separan captada/tierra/inducida/total; `D_anualizacion_escenarios.csv` y `D_resumen_anual_ambos_sentidos.csv` documentan escenarios anuales. Los NaN son faltantes, no cero.

Código: `scripts/03_estimar_demanda_local.py`; configuración con procedencia de cada fijo: `data/templates/config_estimacion.json`. Desde el proyecto: `python scripts/03_estimar_demanda_local.py`; para invalidar cache: `--refresh`; para otra ubicación: `--anexos /ruta/Anexos`. Se requiere pandas/openpyxl. Pruebas sintéticas sin datos reales: `python -m unittest discover -s tests -p test_estimacion_local.py -v`.

Los caches locales seleccionados en `data/interim/estimacion_preliminar/` contienen claves de enlace y están excluidos de Git; no deben compartirse. Las salidas no contienen identificadores individuales, patentes ni coordenadas; algunas celdas agregadas tienen poca muestra y también se excluyeron de Git a la espera de revisión. No se añadieron archivos a Git, no se hizo commit/push ni se envió ningún dato a servicios externos. `fuentes_integridad.csv` comprueba SHA256 antes/después de las fuentes. La reproducción aritmética y la conservación de flujos se comprueban mediante aserciones; no equivalen a validar representatividad.
'''
    (OUT/'INFORME_RESULTADOS.md').write_text(text)
    headline=both.loc[both.variante.eq('media_zarpe_9') & both.segmento.eq('VL') & both.alcance.eq('cuatro_OD_iniciales')]
    rates=headline.pivot(index='temporada',columns='escenario',values='captada_barcaza_dia').reset_index()
    rates=rates[['temporada','bajo','medio','alto']].rename(columns={'bajo':'captada_10pct','medio':'captada_25pct','alto':'captada_50pct'})
    frequency=med.loc[med.alcance.eq('cuatro_OD_iniciales')].pivot(index='temporada',columns='variante',values='captada_barcaza_dia').reset_index()
    refbrief=ref.loc[ref.variante.eq('ARISTO_original') & ref.segmento.ne('PEATON')][['temporada','segmento','flujo_dia']]
    abrief=ann.loc[ann.alcance.eq('cuatro_OD_iniciales')][['variante','invierno','captada_barcaza_ano']]
    (OUT/'RESUMEN_EQUIPO.md').write_text(f'''# Primera estimación para el equipo — 24/10/2026

**Bahía Jara–Levicán: escenarios cuantitativos parciales de desvío desde la barcaza actual.** No son demanda total del proyecto ni predicción frente a la SBO de dos naves.

## 1. Qué entrega ARISTO

Se reprodujeron las fórmulas de TD Transb. Flujos diarios equivalentes originales, ambos sentidos:

{mdtable(refbrief)}

VL conserva la clase de privados de ARISTO, incluyendo bicicletas; CAM conserva su clase de carga, incluyendo furgones/remolques. No son pasajeros ni toneladas. El cálculo original usa conteo acumulado×9/7. ARISTO prescribe en §2.2.6 usar **media por servicio**, y tabla78 identifica dos zarpes medidos por sentido en primavera. Por ello se conservan los originales y se comparan con media por zarpe. La frecuencia12 procede del itinerario de tabla38 y no está comprobada para todo el año; el nueve de las fórmulas tampoco se confirmó como itinerario ordinario de todas las campañas.

## 2. Primera estimación del nuevo cruce

Cuatro OD: Chile Chico↔Coyhaique y Chile Chico↔Río Ibáñez. Solo desvío desde la barcaza; ambos sentidos; unidades de la clase privados/día equivalente. Variante **media por zarpe y nueve semanales**, condicionada a una fila por zarpe en verano/Normal y composición OD muestral representativa por sentido/campaña/clase:

{mdtable(rates)}

Las tasas10/25/50% son sensibilidades exógenas de EVALSO y la solicitud del equipo, no probabilidades calibradas. No se distribuyeron flujos con peatones y vehículos mezclados. La asignación de bicicletas incluidas por ARISTO a OD vehiculares es una aproximación explícita. No deben redondearse a enteros antes de agregar.

Sensibilidad de captación al25%, mismas unidades y cuatro OD:

{mdtable(frequency)}

Se entregan además todos los OD observados en la barcaza y escenarios para el cruce completo como **envolvente**, sin acreditar que todos sean elegibles para Jara–Levicán. Camiones y peatones tienen tablas separadas; hay estratos sin muestra OD y permanecen sin asignar. En particular, camiones CC→Ibáñez y peatones Ibáñez→CC en verano impiden cerrar esos totales locales.

## 3. Anualización ilustrativa

Al25% de captación, cuatro OD, clase privados, unidades/año:

{mdtable(abrief)}

Calendario **supuesto**: verano90 días, primavera91, Normal92 e invierno92; invierno reducido=50% de Normal, igual=100%. No se dispone de medición invernal ni se proyectó crecimiento. La variante tabla169 reconoce que los conteos de mayo del informe difieren del Excel. Estas anualizaciones son escenarios, no demanda anual observada.

## 4. Qué falta y cómo usar el resultado

Tierra, inducción y total del proyecto son **no identificados**, no cero. La SBO incorpora dos naves y mejor frecuencia: no se multiplicó la demanda histórica por dos ni se simuló su respuesta a esa oferta. Faltan itinerario efectivo por temporada y clase, manifiestos, representatividad de zarpes, invierno, ruta terrestre verificada y costos generalizados completos histórico/SBO/proyecto. No se calibró logit ni se inventó elasticidad.

La distribución OD usa entrevistas válidas del cruce dentro de campaña×sentido×clase. Es un supuesto provisional, no ponderación de diseño verificada. Los resúmenes previos de IA no respaldan los números: se usaron los Excel y los informes originales. Las matrices de ARISTO revisadas no contienen los cuatro OD iniciales; además sus fórmulas emplean factores mensuales aunque el informe las describe como diarias. No se usaron para estos flujos locales.

El informe detallado es `INFORME_RESULTADOS.md`; tablas A/B/C/D, fórmulas, caches y fuentes están en esta carpeta. Reproducción desde el proyecto: `python scripts/03_estimar_demanda_local.py`. Código y supuestos están documentados en `docs/estimacion_preliminar.md` y `data/templates/config_estimacion.json`. Los originales permanecen intactos; microdatos y resultados detallados están excluidos de Git. No se hizo push.
''')

def formula_text(v):
    return getattr(v, 'text', v)

def cache_frame(path, sheet, tag, digest, refresh=False, **kwargs):
    """Cache agregado/seleccionado local excluido de Git; invalidado por SHA256."""
    dest = CACHE / f'{tag}.pkl'
    meta = dest.with_suffix('.sha256')
    if not refresh and dest.exists() and meta.exists() and meta.read_text()==digest:
        return pd.read_pickle(dest)
    df = pd.read_excel(path, sheet_name=sheet, **kwargs)
    df.to_pickle(dest)
    meta.write_text(digest)
    return df

def auditar_excel(paths):
    """Fórmula y valor almacenado, sin evaluar Excel ni actualizar vínculos."""
    inventory, links, matrix_checks, factor_rows, td_rows = [], [], [], [], []
    for key in ['eod','conteos','factores','matrices']:
        p = paths[key]
        with ZipFile(p) as z:
            for n in z.namelist():
                if n.startswith('xl/externalLinks/_rels/'):
                    for e in ET.fromstring(z.read(n)):
                        target = unquote(e.attrib.get('Target',''))
                        basename = target.replace('\\','/').rsplit('/',1)[-1]
                        links.append(dict(archivo=p.name,vinculo=n,destino_original=target,
                                          archivo_candidato=basename,existe_nombre_local=any(x.name==basename for x in paths.values()) or (p.parent/basename).exists(),
                                          nota='Ruta original no resuelta automaticamente; candidatos por nombre, sin actualizar Excel'))
        f = load_workbook(p,read_only=True,data_only=False)
        v = load_workbook(p,read_only=True,data_only=True)
        if key=='eod': specs=[('TD Transb',4,110,8),('TD Carr',2,40,8)]
        elif key=='conteos': specs=[('TD',4,160,5),('Transbordadores',1,232,27)]
        elif key=='factores': specs=[('arcos_ptocontrol',1,154,15)]+[(s,1,274,61) for s in f.sheetnames if s.endswith((' VL',' VP'))]
        else: specs=[(s,1,274,61) for s in f.sheetnames]
        for sn,start,end,ncols in specs:
            fr=f[sn].iter_rows(min_row=start,max_row=end,max_col=ncols)
            vr=v[sn].iter_rows(min_row=start,max_row=end,max_col=ncols)
            for rr,(rf,rv) in enumerate(zip(fr,vr),start):
                vals=[c.value for c in rv]
                is_lago=any(es_lago(x) for x in vals[:5] if isinstance(x,str))
                selected=is_lago or (sn.endswith((' VL',' VP')) and rr in [1,2,3,4]) or (sn=='TD Carr' and 'cerro castillo' in norm(vals[0]))
                if selected:
                    for cf,cv in zip(rf,rv):
                        ft=formula_text(cf.value)
                        if isinstance(ft,str) and ft.startswith('='):
                            inventory.append(dict(archivo=p.name,hoja=sn,celda=cf.coordinate,formula=ft,
                                                  valor_almacenado=cv.value,error_excel=cv.data_type=='e'))
                if key=='factores' and sn=='arcos_ptocontrol' and is_lago:
                    factor_rows.append(dict(fila_excel=rr,arco=vals[0],temporada=str(vals[3]).strip(),sentido=sentido(vals[4]),
                        encuestas_VL=vals[5],encuestas_CAM=vals[6],flujo_VL_dia=vals[7],flujo_CAM_dia=vals[8],
                        factor_VL=vals[9],factor_CAM=vals[10],flujo_VL_mes=vals[11],flujo_CAM_mes=vals[12],factor_VL_mes=vals[13],factor_CAM_mes=vals[14]))
                if key=='eod' and sn=='TD Transb' and is_lago:
                    td_rows.append(dict(fila_excel=rr,temporada=str(vals[1]).strip(),sentido=sentido(vals[2]),
                        suma_VL=vals[3],suma_CAM=vals[4],frecuencia_semana=vals[5],VL_dia=vals[6],CAM_dia=vals[7]))
                if sn.endswith((' VL',' VP')) and rr>=3 and numero(vals[0]) and numero(vals[1]):
                    den=sum(x for x in vals[5:] if numero(x))
                    expected=vals[4]/den if numero(vals[4]) and den else 0
                    matrix_checks.append(dict(archivo=p.name,hoja=sn,fila_excel=rr,origen=vals[0],destino=vals[1],
                        resultado_cache=vals[3],encuestas=vals[4],denominador_cache=den,reconstruido=expected,
                        coincide=numero(vals[3]) and math.isclose(vals[3],expected,rel_tol=1e-9,abs_tol=1e-9),
                        par_inicial=(int(vals[0]),int(vals[1])) in [(39,31),(31,39),(39,40),(40,39)],
                        unidad_formula='vehiculos_mes_si_factor_fila1_es_encuestas_sobre_flujo_mensual'))
        f.close();v.close()
    return pd.DataFrame(inventory),pd.DataFrame(links),pd.DataFrame(matrix_checks),pd.DataFrame(factor_rows),pd.DataFrame(td_rows)

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--anexos',type=Path,default=ROOT.parent)
    parser.add_argument('--config',type=Path,default=ROOT/'data/templates/config_estimacion.json')
    parser.add_argument('--refresh',action='store_true')
    args=parser.parse_args()
    OUT.mkdir(parents=True,exist_ok=True);CACHE.mkdir(parents=True,exist_ok=True)
    config=json.loads(args.config.read_text())
    paths={k:args.anexos/v for k,v in config['fuentes'].items()}
    paths['itinerario_dbf']=args.anexos/'SIG/Coberturas/ARISTO/Transbordadores.dbf'
    hashes={k:hashlib.sha256(p.read_bytes()).hexdigest() for k,p in paths.items()}
    print('Leyendo solo fuentes seleccionadas y conservando cache local por hash...',flush=True)
    cols1=['ID 2','Versión','Validada','Duplicada','Interregional','Lugar encuestado','Sentido Viaje','Día','Encuesta VL','Encuesta Cam','Propósito de viaje','Lugar de Residencia']
    a=cache_frame(paths['eod'],'BBDD-1','eod_original_seleccion',hashes['eod'],args.refresh,
                  usecols=lambda c:c in cols1 or ('transbordador' in str(c).lower() and ('usad' in str(c).lower() or 'cuál' in str(c).lower())))
    b=cache_frame(paths['eod'],'BBDD-2 (solo válidas)','eod_valida_seleccion',hashes['eod'],args.refresh,
                  usecols=['ID 2','Nombre origen','Nombre destino','Tipo usuario'])
    assert a['ID 2'].notna().all() and b['ID 2'].notna().all()
    assert a['ID 2'].is_unique and b['ID 2'].is_unique
    assert set(b['ID 2'])==set(a.loc[a.Validada.eq(0),'ID 2'])
    d=a.merge(b,on='ID 2',validate='one_to_one')
    d['temporada']=d['Versión'].str.strip()
    d['interceptado']=d['Lugar encuestado'].map(es_lago)
    fields=[c for c in a if 'transbordador' in c.lower()]
    d['declarado']=d[fields].apply(lambda s:s.map(es_lago)).any(axis=1)
    d['ruta']='desconocida'
    d.loc[d.declarado,'ruta']='lago_declarado_usado_o_por_usar'
    d.loc[d.interceptado,'ruta']='lago_confirmado_interceptacion'
    d['segmento']=d['Tipo usuario'].map(config['segmentos']).fillna('OTRO_PASAJERO')
    pairs={tuple(x) for x in config['pares_iniciales']}
    d['par_inicial']=[(o,t) in pairs for o,t in zip(d['Nombre origen'],d['Nombre destino'])]
    # Mostrar declaraciones Ninguno, sin convertirlas en itinerario terrestre verificado.
    d['declara_ninguno']=d[fields].apply(lambda s:s.map(norm).eq('ninguno')).any(axis=1)
    save(d.groupby(['Nombre origen','Nombre destino','temporada','segmento','ruta','par_inicial'],dropna=False).size().reset_index(name='encuestas'),'muestra_od_ruta')
    save(d.loc[d.par_inicial & ~d.interceptado].groupby(['temporada','Lugar encuestado','ruta','segmento','declara_ninguno'],dropna=False).size().reset_index(name='encuestas'),'candidatos_fuera_barcaza_NO_flujo_terrestre')
    lake=d.loc[d.interceptado].copy()
    lake['sentido']=lake['Sentido Viaje'].map(sentido)
    nk=['temporada','sentido','segmento']
    samp=lake.groupby(nk+['Nombre origen','Nombre destino','par_inicial'],dropna=False).size().reset_index(name='n_od')
    samp['n_estrato']=samp.groupby(nk).n_od.transform('sum')
    samp['participacion_muestral']=samp.n_od/samp.n_estrato
    save(samp,'participaciones_od_barcaza')
    save(lake.groupby(['temporada','sentido','Tipo usuario','segmento'],dropna=False).agg(encuestas=('ID 2','size'),bandera_VL=('Encuesta VL','sum'),bandera_CAM=('Encuesta Cam','sum')).reset_index(),'control_clasificacion_usuarios')
    count=cache_frame(paths['conteos'],'Transbordadores','conteos',hashes['conteos'],args.refresh)
    count['fila_excel']=count.index+2
    count=count.loc[count.Ruta.map(es_lago)].copy()
    count['temporada']=count.Temporada.str.strip()
    count['sentido']=count['Sentido de viaje'].map(sentido)
    save(count,'conteos_originales_cruce')
    roads=cache_frame(paths['conteos'],'Carretera','carretera',hashes['conteos'],args.refresh)
    roads=roads.loc[roads.PC.eq('ODC-3')].copy()
    rsum=roads.groupby(['PC','DESCRIPCIÓN','Temporada','SENTIDO']).agg(intervalos=('CUARTO','size'),dias_etiquetados=('DÍA','nunique'),privados_observados=('V.Privado','sum'),camiones_observados=('Camión','sum')).reset_index()
    rsum['unidad']='conteos_acumulados_jornada_no_24h'
    rsum['estado']='observado_corredor_no_OD_local_ni_ruta_exclusivamente_terrestre'
    save(rsum,'carretera_Cerro_Castillo_NO_demanda_local')
    rc=cache_frame(paths['eod'],'TD Carr','td_carretera',hashes['eod'],args.refresh,header=1,nrows=35)
    save(rc.loc[rc['DESCRIPCIÓN'].map(norm).str.contains('cerro castillo')],'TD_Carr_Cerro_Castillo_original')
    service=dbf_carrera(paths['itinerario_dbf'])
    (OUT/'itinerario_SIG_original.json').write_text(json.dumps(service,ensure_ascii=False,indent=2))
    calendar=[]
    for direction,values in config['frecuencia_tabla38'].items():
        for day,n in zip(['lunes','martes','miercoles','jueves','viernes','sabado','domingo'],values):
            calendar.append(dict(sentido=direction,dia=day,zarpes_programados=n,fuente='ARISTO tabla38 p112',periodo='catastro_primavera_no_bitacora_efectiva'))
    save(pd.DataFrame(calendar),'itinerario_tabla38')
    print('Microdatos agregados. Extrayendo formulas y caches de la cadena de expansion...',flush=True)
    inv,links,mat,factors,td=auditar_excel(paths)
    for frame,name in [(inv,'formulas_y_cache'),(links,'vinculos_externos'),(mat,'reproduccion_matrices'),(factors,'factores_arcos_originales'),(td,'TD_Transb_original')]:save(frame,name)
    linked=[]
    for _,r in factors.iterrows():
        t=td.loc[td.temporada.eq(r.temporada)&td.sentido.eq(r.sentido)].iloc[0]
        for seg in ['VL','CAM']:
            assert math.isclose(r[f'flujo_{seg}_dia'],t[f'{seg}_dia'],rel_tol=1e-9)
            n=int(samp.loc[samp.temporada.eq(r.temporada)&samp.sentido.eq(r.sentido)&samp.segmento.eq(seg),'n_od'].sum())
            linked.append(dict(temporada=r.temporada,sentido=r.sentido,segmento=seg,encuestas_factor_ARISTO=r[f'encuestas_{seg}'],
                encuestas_distribucion_local=n,diferencia=r[f'encuestas_{seg}']-n,
                flujo_cache_arcos=r[f'flujo_{seg}_dia'],flujo_TD_Transb=t[f'{seg}_dia'],coinciden_flujos=True))
    save(pd.DataFrame(linked),'control_enlace_TD_arcos_y_muestra')
    recon=[]
    for _,r in factors.iterrows():
        for seg in ['VL','CAM']:
            q=r[f'flujo_{seg}_dia'];n=r[f'encuestas_{seg}']
            for what,expected,actual in [('flujo_mes',q*30.5,r[f'flujo_{seg}_mes']),('factor',n/q if n and q else float('nan'),r[f'factor_{seg}']),('factor_mes',n/(q*30.5) if n and q else float('nan'),r[f'factor_{seg}_mes'])]:
                recon.append(dict(temporada=r.temporada,sentido=r.sentido,segmento=seg,calculo=what,reconstruido=expected,cache=actual,
                                  coincide=(pd.isna(expected) and pd.isna(actual)) or (numero(expected) and numero(actual) and math.isclose(expected,actual,rel_tol=1e-9))))
    save(pd.DataFrame(recon),'reproduccion_factores')
    comparison=[]
    # Mantener el calculo original y dos alternativas con divisor por exposicion.
    for (season,direction),g in count.groupby(['temporada','sentido']):
        t=td.loc[td.temporada.eq(season)&td.sentido.eq(direction)].iloc[0]
        for seg,col in [('VL','Vehículos Privados'),('CAM','Camiones'),('PEATON','Peatones')]:
            observed=g[col].sum(min_count=1)
            for variant in config['variantes']:
                basis=observed
                if variant=='ARISTO_original':
                    q=t[f'{seg}_dia'] if seg!='PEATON' else float('nan')
                    freq=t.frecuencia_semana;exposure=float('nan')
                else:
                    freq=t.frecuencia_semana if variant=='media_zarpe_9' else sum(config['frecuencia_tabla38'][direction])
                    exposure=len(g)*config['zarpes_por_fila'][season]['valor']
                    if variant=='media_zarpe_12_tabla169' and season=='Normal' and seg!='PEATON':
                        basis=config['tabla169_normal'][direction][seg]
                    q=basis/exposure*freq/7
                original=t[f'{seg}_dia'] if seg!='PEATON' else float('nan')
                comparison.append(dict(variante=variant,temporada=season,sentido=direction,segmento=seg,unidad='personas_peatones_sin_bicicletas_dia' if seg=='PEATON' else ('unidades_clase_privados_ARISTO_dia_incluye_bicicletas' if seg=='VL' else 'vehiculos_carga_ARISTO_dia'),
                    conteo_acumulado_Excel=observed,conteo_base_calculo=basis,filas_conteo=len(g),zarpes_usados_como_divisor=exposure,frecuencia_semana=freq,
                    flujo_dia=q,flujo_semana=q*7,ARISTO_dia=original,
                    diferencia_frente_ARISTO=q-original,estado='calculado_ARISTO_frecuencia_no_conciliada' if variant=='ARISTO_original' else 'calculado_condicionado',
                    representatividad=config['zarpes_por_fila'][season]['estado'],fuente='TD Transb y Transbordadores; config; ARISTO 2.2.6 y tablas38/78/107/138/169'))
                if seg!='PEATON' and variant=='ARISTO_original':assert math.isclose(q,observed*freq/7,rel_tol=1e-9)
    comp=pd.DataFrame(comparison);save(comp,'B_comparacion_expansion')
    assert pd.DataFrame(recon).coincide.all()
    local=[]
    for _,r in comp.loc[comp.flujo_dia.notna()].iterrows():
        selected=samp.loc[samp.temporada.eq(r.temporada)&samp.sentido.eq(r.sentido)&samp.segmento.eq(r.segmento)]
        if selected.empty:
            local.append(dict(**r.to_dict(),origen='OD_no_identificado',destino='OD_no_identificado',par_inicial=False,n_od=0,n_estrato=0,participacion=float('nan'),flujo_od_dia=float('nan'),flujo_sin_asignar_dia=r.flujo_dia))
        for _,s in selected.iterrows():
            local.append(dict(**r.to_dict(),origen=s['Nombre origen'],destino=s['Nombre destino'],par_inicial=s.par_inicial,n_od=s.n_od,n_estrato=s.n_estrato,
                participacion=s.participacion_muestral,flujo_od_dia=repartir(r.flujo_dia,s.n_od,s.n_estrato),flujo_sin_asignar_dia=0))
    local=pd.DataFrame(local)
    local['metodo_od']='supuesto_proporcion_muestral_dentro_de_temporada_sentido_y_segmento; solo_interceptadas_en_cruce'
    local['estado_od']='calculado_con_distribucion_OD_supuesta'
    save(local,'A_demanda_referencia_od')
    # Comprobacion de conservacion incluyendo estratos sin muestra.
    balance=local.groupby(['variante','temporada','sentido','segmento']).agg(asignado=('flujo_od_dia',lambda s:s.sum(min_count=1)),sin_asignar=('flujo_sin_asignar_dia','sum'),total=('flujo_dia','first')).reset_index()
    balance['diferencia']=balance.asignado.fillna(0)+balance.sin_asignar-balance.total
    assert balance.diferencia.abs().max()<1e-9
    save(balance,'control_conservacion_flujos')
    projects=[]
    for _,r in local.iterrows():
        for scenario,rate in config['captacion'].items():
            projects.append(dict(variante=r.variante,temporada=r.temporada,sentido=r.sentido,segmento=r.segmento,unidad=r.unidad,origen=r.origen,destino=r.destino,
                par_inicial=r.par_inicial,escenario=scenario,captacion_supuesta=rate,base_od_dia=r.flujo_od_dia,
                captada_barcaza_dia=r.flujo_od_dia*rate,tierra_dia=float('nan'),inducida_dia=float('nan'),total_proyecto_dia=float('nan'),
                estado='escenario_parcial_desviacion_barcaza_historica_sin_respuesta_SBO'))
    project=pd.DataFrame(projects);save(project,'C_captacion_por_od')
    summary=[]
    # Envolvente del cruce != mercado elegible. No perder estratos sin muestra al resumir.
    for _,r in comp.loc[comp.flujo_dia.notna()].iterrows():
        for scope in ['cruce_completo_envolvente','cuatro_OD_iniciales']:
            z=local.loc[local.variante.eq(r.variante)&local.temporada.eq(r.temporada)&local.sentido.eq(r.sentido)&local.segmento.eq(r.segmento)]
            q=r.flujo_dia if scope=='cruce_completo_envolvente' else (z.loc[z.par_inicial,'flujo_od_dia'].sum() if z.n_estrato.max()>0 else float('nan'))
            for scenario,rate in config['captacion'].items():
                summary.append(dict(variante=r.variante,temporada=r.temporada,sentido=r.sentido,segmento=r.segmento,unidad=r.unidad,alcance=scope,
                    escenario=scenario,captacion_supuesta=rate,referencia_dia=q,captada_barcaza_dia=q*rate,tierra_dia=float('nan'),inducida_dia=float('nan'),total_proyecto_dia=float('nan')))
    summary=pd.DataFrame(summary);save(summary,'C_resumen_escenarios')
    sk=['variante','temporada','segmento','unidad','alcance','escenario','captacion_supuesta']
    both=summary.groupby(sk).agg(referencia_dia=('referencia_dia',lambda s:s.sum(min_count=2)),captada_barcaza_dia=('captada_barcaza_dia',lambda s:s.sum(min_count=2))).reset_index()
    both['tierra_dia']=float('nan');both['inducida_dia']=float('nan');both['total_proyecto_dia']=float('nan')
    save(both,'C_resumen_ambos_sentidos')
    annual=[]
    for keys,g in summary.groupby(['variante','sentido','segmento','alcance','escenario','unidad']):
        flows=dict(zip(g.temporada,g.captada_barcaza_dia))
        for winter,ratio in config['invierno_sobre_normal'].items():
            annual.append(dict(zip(['variante','sentido','segmento','alcance','escenario','unidad_diaria'],keys),
                invierno= winter,razon_invierno_normal=ratio,captada_barcaza_ano=anualizar(flows,config['calendario_anual'],ratio),
                tierra_ano=float('nan'),inducida_ano=float('nan'),total_proyecto_ano=float('nan'),
                unidad_anual='personas_peatones_ano' if keys[2]=='PEATON' else 'vehiculos_ano',estado='escenario_anual_calendario_y_perfil_invernal_supuestos'))
    annual=pd.DataFrame(annual);save(annual,'D_anualizacion_escenarios')
    annual['referencia_barcaza_ano']=[v/config['captacion'][s] for v,s in zip(annual.captada_barcaza_ano,annual.escenario)]
    save(annual,'D_anualizacion_escenarios')
    ak=['variante','segmento','alcance','escenario','invierno','razon_invierno_normal','unidad_anual']
    annual_both=annual.groupby(ak).agg(captada_barcaza_ano=('captada_barcaza_ano',lambda s:s.sum(min_count=2))).reset_index()
    annual_both['tierra_ano']=float('nan');annual_both['inducida_ano']=float('nan');annual_both['total_proyecto_ano']=float('nan')
    save(annual_both,'D_resumen_anual_ambos_sentidos')
    # Otros OD del cruce: presentados como candidatos, no validados por accesos.
    save(samp.loc[~samp.par_inicial],'otros_OD_candidatos_sin_elegibilidad_confirmada')
    assert sum(config['calendario_anual'].values())==365
    manifest=[]
    for k,p in paths.items():
        after=hashlib.sha256(p.read_bytes()).hexdigest();assert hashes[k]==after
        manifest.append(dict(fuente=str(p.relative_to(args.anexos)),sha256=after,original_sin_cambios=True))
    save(pd.DataFrame(manifest),'fuentes_integridad')
    (OUT/'config_utilizada.json').write_text(json.dumps(config,ensure_ascii=False,indent=2))
    controls={'originales':len(a),'validas':len(b),'interceptadas_cruce':len(lake),'pares_iniciales':int(d.par_inicial.sum()),
              'validas_bandera_duplicada':int(d.Duplicada.eq(1).sum()),'matrices_pares_iniciales_presentes':int(mat.par_inicial.sum()),
              'formulas_matriz_reproducidas':int(mat.coincide.sum()),'formulas_matriz_total':len(mat),
              'factores_reproducidos':len(recon),'balance_max_error':float(balance.diferencia.abs().max())}
    (OUT/'controles.json').write_text(json.dumps(controls,ensure_ascii=False,indent=2))
    reporte(config,comp,both,annual_both,controls,samp)
    print(json.dumps(controls,ensure_ascii=False,indent=2),flush=True)
    print('Flujos originales ARISTO por dia, ambos sentidos (sin peatones):',flush=True)
    print(comp.loc[comp.variante.eq('ARISTO_original')].groupby(['temporada','segmento']).flujo_dia.sum(min_count=1).to_string(),flush=True)
    print('Resultados guardados en',OUT,flush=True)

if __name__=='__main__':main()
