"""Matrices OD por ruta: barcaza primero, expansion terrestre parcial y Argentina.

No ejecuta SBO, ocupaciones, captaciones, induccion ni logit.
"""
from pathlib import Path
import sys,json,math,hashlib,importlib.util
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'src'))
from demanda.capacidad import documento_legible
spec=importlib.util.spec_from_file_location('prev',ROOT/'scripts/03_estimar_demanda_local.py')
prev=importlib.util.module_from_spec(spec);spec.loader.exec_module(prev)
OUT=ROOT/'outputs/matrices_od_referencia';CACHE=ROOT/'data/interim/estimacion_preliminar'

def save(d,n):d.to_csv(OUT/f'{n}.csv',index=False,encoding='utf-8-sig',na_rep='NaN')


def resolver_lugar(row,side,cfg):
    original=row['Nombre origen' if side=='origen' else 'Nombre destino']
    if original!='Río Ibáñez':return original,'zona_o_comuna_BBDD2'
    fields=['Búsqueda Origen','Dirección Origen'] if side=='origen' else ['Búsqueda destino','Dirección destino']
    txt=' '.join(prev.norm(row.get(c)) for c in fields)
    found=[name for name,tokens in cfg['localidades_rio_ibanez'].items() if any(t in txt for t in tokens)]
    if len(found)==1:return found[0],'localidad_por_mencion_textual_EOD_no_poligono'
    return 'Río Ibáñez','comuna_o_zona_sin_localidad_univoca'


def clasificar(row,fs,bs,cfg,rules):
    fv=[prev.norm(row.get(c)) for c in fs];bv=[prev.norm(row.get(c)) for c in bs]
    neg=set(rules['sin_transbordador_explicito'])
    fpos=[v for v in fv if v and v not in neg]
    nof=any(v in neg for v in fv) and not fpos
    bpos=[v for v in bv if v and v not in cfg['sin_fronteras']]
    nob=any(v in cfg['sin_fronteras'] for v in bv) and not bpos
    steps=sorted({t for t in rules['fronteras_argentina'] if any(t in v for v in bpos)})
    place=prev.norm(row['Lugar encuestado'])
    border_control=any(t in place for t in rules['fronteras_argentina'])
    land_control='carretera' in prev.norm(row['Lugar de Encuesta']) or border_control
    if prev.es_lago(place):route,grade,evidence='barcaza_actual','interceptacion_confirmada','encuestado_en_cruce_actual'
    elif nof and nob and land_control and not border_control:
        route,grade,evidence='terrestre_chile','declaracion_expresa','ningun_transbordador_y_ningun_paso_fronterizo; control_terrestre_Chile'
    elif nof and steps:
        route='terrestre_argentina'
        grade='probable_varios_pasos_declarados_sin_trazado_completo' if len(steps)>=2 else 'probable_itinerario_incompleto_un_paso'
        evidence='sin_transbordador; pasos_declarados; interceptacion_frontera_no_usada_como_prueba_de_recorrido_completo'
    elif nof and land_control:
        route,grade,evidence='terrestre_pais_desconocido','declaracion_parcial','sin_transbordador; pais_no_demostrado'
    else:
        route,grade,evidence='ruta_desconocida','sin_itinerario_completo','vacios_o_declaraciones_no_concluyentes_o_mixtas'
    return dict(ruta=route,grado_evidencia=grade,evidencia=evidence,pasos_fronterizos='; '.join(steps) or ('Ninguno_declarado' if nob else 'no_identificado'),
                control_fronterizo_observado=border_control,lago_declarado=any(prev.es_lago(v) for v in fpos),
                ningun_transbordador_explicito=nof,ninguna_frontera_explicita=nob)


def agrupar_encuestas(d):
    keys=['estacion','origen','destino','sentido','segmento','ruta','grado_evidencia','pasos_fronterizos','evidencia','ambito','nivel_origen','nivel_destino']
    return d.groupby(keys,dropna=False).size().reset_index(name='encuestas_validas')


def metadatos(r):
    units={'VL':'vehiculos_livianos','CAM':'vehiculos_carga','PEATON':'personas_peatones','MOTO':'motocicletas','BICI':'bicicletas',
           'OTRO_PASAJERO':'pasajero_individual_sin_unidad_fisica_asignable'}
    return dict(estacion=r.estacion,origen=r.origen,destino=r.destino,sentido=r.sentido,
                segmento=r.segmento,unidad=units.get(r.segmento,'unidad_fisica_no_clasificada'),
                nivel_origen=r.nivel_origen,nivel_destino=r.nivel_destino,ambito=r.ambito)


def main():
    OUT.mkdir(exist_ok=True,parents=True);CACHE.mkdir(exist_ok=True,parents=True)
    cfg=json.loads((ROOT/'data/templates/config_matrices_od.json').read_text())
    rules=json.loads((ROOT/'data/templates/config_rutas_terrestres.json').read_text())
    # El contraste CITRA es opcional; sus tablas originales no se distribuyen.
    citra_path=ROOT/'data/templates/citra_matriz_publicada.json'
    if not citra_path.exists():citra_path=ROOT/'data/templates/citra_matriz_publicada.example.json'
    citra_cfg=json.loads(citra_path.read_text())
    eod=ROOT.parent/'Cap 6/BBDD EOD/Base de datos agregada.xlsx'
    counts=ROOT.parent/'Cap 6/Mediciones de flujo/Consolidado Flujos.xlsx'
    points=ROOT.parent/'Cap 8/EOD_Puntos_Pasada_Final.xlsx'
    arcs=ROOT.parent/'Cap 8/BBDD_para_Pijas_actualizada_EOD_Final.xlsx'
    deod=hashlib.sha256(eod.read_bytes()).hexdigest();dcounts=hashlib.sha256(counts.read_bytes()).hexdigest()
    dpoints=hashlib.sha256(points.read_bytes()).hexdigest();darcs=hashlib.sha256(arcs.read_bytes()).hexdigest()
    cols1=['ID 2','Versión','Validada','Lugar encuestado','Sentido Viaje','Día']
    a=prev.cache_frame(eod,'BBDD-1','eod_original_seleccion',deod,False,
        usecols=lambda c:c in cols1 or ('transbordador' in str(c).lower() and ('usad' in str(c).lower() or 'cuál' in str(c).lower())))
    b=prev.cache_frame(eod,'BBDD-2 (solo válidas)','eod_valida_seleccion',deod,False,
        usecols=['ID 2','Nombre origen','Nombre destino','Tipo usuario'])
    more=prev.cache_frame(eod,'BBDD-1','eod_ubicacion_fronteras_matrices',deod,False,
        usecols=lambda c:c in ['ID 2','Lugar de Encuesta'] or 'fronterizo' in str(c).lower())
    geo=prev.cache_frame(eod,'BBDD-1','eod_direcciones_horario_matrices',deod,False,
        usecols=['ID 2','Búsqueda Origen','Dirección Origen','Búsqueda destino','Dirección destino','Hora'])
    d=a.merge(b,on='ID 2',validate='one_to_one').merge(more,on='ID 2',validate='one_to_one').merge(geo,on='ID 2',validate='one_to_one')
    d['campana']=d['Versión'].str.strip();d['estacion']=d.campana.map(cfg['estacion_campana'])
    d['segmento']=d['Tipo usuario'].map(cfg['segmentos']).fillna('OTRO_PASAJERO')
    for side in ['origen','destino']:
        z=d.apply(lambda r:resolver_lugar(r,side,cfg),axis=1)
        d[side]=z.map(lambda v:v[0]);d['nivel_'+side]=z.map(lambda v:v[1])
    d['sentido']=d.origen+' → '+d.destino
    fs=[c for c in a if 'transbordador' in c.lower()];bs=[c for c in more if 'fronterizo' in c.lower()]
    flags=d.apply(lambda r:pd.Series(clasificar(r,fs,bs,cfg,rules)),axis=1);d=pd.concat([d,flags],axis=1)
    core={tuple(p) for p in cfg['pares_nucleo']}
    d['nucleo']=[(o,t) in core for o,t in zip(d['Nombre origen'],d['Nombre destino'])]
    d['ambito']=d.nucleo.map({True:'nucleo_Chile_Chico_ribera_norte',False:'otros_OD_observados_cruce_no_elegibilidad_confirmada'})
    rp=prev.cache_frame(points,'Puntos por OD (Ordenados)','rutas_puntos',dpoints,False,
        usecols=['ID 2','Ruta OD','OD-Nodos del recorrido'])
    ra=prev.cache_frame(arcs,'Arcos por OD','rutas_arcos',darcs,False,usecols=['ID 2','Nodos_Pasada'])
    coded=d[['ID 2','campana','estacion','origen','destino','sentido','segmento','nucleo']].merge(rp,on='ID 2',how='left',validate='one_to_one').merge(ra,on='ID 2',how='left',validate='one_to_one')
    coded=coded.loc[coded.nucleo].groupby(['estacion','origen','destino','sentido','segmento'],dropna=False).agg(
        encuestas_validas=('ID 2','size'),registros_con_ruta_codificada=('Ruta OD',lambda s:s.notna().sum()),
        registros_con_nodos_ruta=('OD-Nodos del recorrido',lambda s:s.notna().sum()),
        registros_con_arcos=('Nodos_Pasada',lambda s:s.notna().sum())).reset_index()
    coded['alcance']='Rutas/codigos de Cap8 no se enlazan a microdatos del OD nucleo si el conteo es cero; no inferir ruta desde su ausencia'
    # Barcaza: mantener todo el universo de interceptacion para conservacion.
    lake=d.loc[d.ruta.eq('barcaza_actual')].copy()
    lake['sentido_cruce']=lake['Sentido Viaje'].map(prev.sentido)
    # Reconstruir el flujo del cruce directamente desde el Excel original, sin ejecutar el módulo SBO.
    ferry=prev.cache_frame(counts,'Transbordadores','conteos',dcounts,False)
    ferry=ferry.loc[ferry.Ruta.map(prev.es_lago)].copy()
    ferry['campana']=ferry.Temporada.str.strip();ferry['sentido']=ferry['Sentido de viaje'].map(prev.sentido)
    ferry['VL']=ferry[['vehículos livianos','Camioneta','Minivan','Van','Casa Rodante']].sum(axis=1,min_count=1)
    ferry['CAM']=ferry['Camiones'];ferry['PEATON']=ferry['Peatones'];ferry['MOTO']=ferry['Motos'];ferry['BICI']=ferry['Bicicletas']
    means=ferry.groupby(['campana','sentido'])[['VL','CAM','PEATON','MOTO','BICI']].agg(lambda s:s.mean(skipna=False)).reset_index()
    row_meta=ferry.groupby(['campana','sentido']).agg(filas_contadas=('FECHA ','size'),fechas_conteo=('FECHA ',lambda s:';'.join(sorted(pd.to_datetime(s).dt.strftime('%Y-%m-%d').unique())))).reset_index()
    means=means.merge(row_meta,on=['campana','sentido'],validate='one_to_one')
    representative=[]
    camp_by_station={'Verano':'Verano','Otoño':'Normal','Primavera':'Primavera'}
    for station,month in cfg['mes_representativo_barcaza'].items():
        campaign=camp_by_station[station]
        frequency=cfg['frecuencia_ordinaria_zarpes_semanales_por_mes'][str(month)]
        z=means.loc[means.campana.eq(campaign)]
        for rr in z.itertuples():
            for seg in ['VL','CAM','PEATON','MOTO','BICI']:
                per_sailing=getattr(rr,seg)
                representative.append(dict(estacion=station,mes=month,segmento=seg,sentido=rr.sentido,
                    campana_fuente=campaign,media_por_zarpe=per_sailing,frecuencia_ordinaria_semana=frequency,
                    filas_contadas_asimiladas_a_zarpes=rr.filas_contadas,fechas_conteo=rr.fechas_conteo,
                    flujo_dia=per_sailing*frequency/7 if pd.notna(per_sailing) else math.nan,
                    unidad='personas_peatones' if seg=='PEATON' else 'unidades_clase_'+seg,
                    estado='estimacion_condicional_atendida' if pd.notna(per_sailing) else 'no_identificado',
                    metodo='media_por_fila_zarpe*frecuencia_ordinaria/7; transferencia_solo_perfil_campana'))
    representative=pd.DataFrame(representative)
    refs=representative.copy()
    # Contrastar fórmula original TD Transb con conteos por fila y cálculo actualizado.
    wf=load_workbook(eod,read_only=True,data_only=False,keep_links=False)
    wv=load_workbook(eod,read_only=True,data_only=True,keep_links=False)
    sf=wf['TD Transb'];sv=wv['TD Transb'];tdrows=[]
    for row in range(20,26):
        route,season,direction=sf.cell(row,1).value,sf.cell(row,2).value,sf.cell(row,3).value
        if not prev.es_lago(route):continue
        campaign=str(season).strip();control=prev.sentido(direction)
        g=ferry.loc[ferry.campana.eq(campaign)&ferry.sentido.eq(control)]
        month=cfg['mes_representativo_barcaza'][{'Verano':'Verano','Normal':'Otoño','Primavera':'Primavera'}[campaign]]
        freq=cfg['frecuencia_ordinaria_zarpes_semanales_por_mes'][str(month)]
        private=g[['VL','MOTO','BICI']].sum(axis=1,min_count=1)
        private_sum=private.sum(min_count=1);cam_sum=g.CAM.sum(min_count=1);n_rows=len(g)
        d_orig,e_orig,f_orig=sf.cell(row,4).value,sf.cell(row,5).value,sf.cell(row,6).value
        tdrows.append(dict(fila_TD_Transb=row,campana=campaign,sentido=control,
            formula_G=sf.cell(row,7).value,valor_G_cache=sv.cell(row,7).value,
            formula_H=sf.cell(row,8).value,valor_H_cache=sv.cell(row,8).value,
            D_privados_ARISTO=d_orig,E_camiones_ARISTO=e_orig,F_frecuencia_original=f_orig,
            filas_conteo_observadas=n_rows,conteo_privados_reconstruido=private_sum,conteo_CAM_reconstruido=cam_sum,
            conteo_D_coincide_sum_filas=bool(pd.notna(d_orig) and pd.notna(private_sum) and math.isclose(float(d_orig),float(private_sum))),
            conteo_E_coincide_sum_filas=bool(pd.notna(e_orig) and pd.notna(cam_sum) and math.isclose(float(e_orig),float(cam_sum))),
            frecuencia_ordinaria_actual=freq,
            privado_por_fila=private.mean(skipna=False) if n_rows else math.nan,
            CAM_por_fila=g.CAM.mean(skipna=False) if n_rows else math.nan,
            flujo_privados_actual=private.mean(skipna=False)*freq/7 if n_rows else math.nan,
            flujo_CAM_actual=g.CAM.mean(skipna=False)*freq/7 if n_rows else math.nan,
            nota='private ARISTO incluye VL+MOTO+BICI; matriz los separa; media por fila aplicada porque formula metodologica es media por zarpe*frecuencia semanal/7'))
    wf.close();wv.close()
    tdcompare=pd.DataFrame(tdrows)
    save(tdcompare,'comparacion_TD_Transb_original_vs_reconstruccion')
    bar=[];balance=[];unassigned=[]
    for rr in representative.itertuples():
        z=lake.loc[lake.campana.eq(rr.campana_fuente)&lake.sentido_cruce.eq(rr.sentido)&lake.segmento.eq(rr.segmento)]
        n=len(z);alloc=0
        if n:
            grouped=z.groupby(['origen','destino','sentido','nivel_origen','nivel_destino','ambito'],dropna=False).size().reset_index(name='n')
            for r in grouped.itertuples():
                q=rr.flujo_dia*r.n/n;alloc+=q if math.isfinite(q) else 0
                bar.append(dict(estacion=rr.estacion,origen=r.origen,destino=r.destino,sentido=r.sentido,segmento=rr.segmento,
                    ruta='barcaza_actual',unidad=rr.unidad,nivel_origen=r.nivel_origen,nivel_destino=r.nivel_destino,ambito=r.ambito,
                    sentido_control=rr.sentido,encuestas_validas=r.n,n_estrato=n,factor_expansion=rr.flujo_dia/n,
                    demanda_expandida_viajes_diarios_promedio=q,demanda_diaria_24h_equivalente=q,
                    procedimiento='flujo_cruce_media_por_zarpe*frecuencia_ordinaria/7 * n_OD/n_estrato',
                    periodo_representado=cfg['periodo_barcaza'][rr.campana_fuente],estacion_completa=rr.estacion!='Otoño',
                    fuente='ARISTO Cap6 BBDD1+BBDD2+Transbordadores; frecuencias ordinarias usuario; paso04 sin ejecutar SBO',
                    validacion='condicional_proporcion_muestral; no_diseno_certificado',
                    limitaciones='una_fila_por_zarpe_en_verano_mayo_supuesto; conteos_sin_indicador_carga_peligrosa; composicion_respuestas_representativa; no_todos_los_pasajeros',
                    cobertura_horas_dia='promedio_semanal/7',grado_evidencia='interceptacion_confirmada'))
        residual=rr.flujo_dia if not n else (rr.flujo_dia-alloc if math.isfinite(rr.flujo_dia) else math.nan)
        balance.append(dict(estacion=rr.estacion,sentido_control=rr.sentido,segmento=rr.segmento,referencia_dia=rr.flujo_dia,
                            asignado_OD=alloc if n and math.isfinite(rr.flujo_dia) else math.nan,sin_asignar_dia=residual,n_estrato=n))
        if not n or not math.isfinite(rr.flujo_dia):
            unassigned.append(dict(estacion=rr.estacion,sentido_control=rr.sentido,segmento=rr.segmento,flujo_sin_asignar_dia=residual,
                                   motivo='sin_muestra_compatible' if not n else 'conteo_componente_no_disponible'))
    # Otros pasajeros de la barcaza sin conteo de personas compatible.
    for r in agrupar_encuestas(lake.loc[lake.segmento.eq('OTRO_PASAJERO')]).itertuples():
        bar.append(dict(**metadatos(r),ruta='barcaza_actual',encuestas_validas=r.encuestas_validas,factor_expansion=math.nan,
                    demanda_expandida_viajes_diarios_promedio=math.nan,demanda_diaria_24h_equivalente=math.nan,
                    validacion='sin_conteo_personas_bus_van_transfer_compatible',limitaciones='no_representan_vehiculos',
                    periodo_representado='campana_observada',estacion_completa=r.estacion!='Otoño'))
    bar=pd.DataFrame(bar);bal=pd.DataFrame(balance)
    save(bar,'matriz_od_barcaza');save(bal,'conservacion_flujos_barcaza');save(pd.DataFrame(unassigned),'flujos_barcaza_sin_OD')
    # Guardar inmediatamente el primer producto antes de la investigacion terrestre.
    print('Matriz de barcaza generada:',len(bar),'filas; continuando con control terrestre.',flush=True)
    # Tierra: un solo control que observa el corredor por Chile, sin sumar controles en serie.
    road=prev.cache_frame(counts,'Carretera','carretera',dcounts,False);road=road.loc[road.PC.eq('ODC-3')].copy()
    road['VL']=road[cfg['conteos_VL_Cerro']].sum(axis=1,min_count=1)
    road['CAM']=road['Camión'];road['MOTO']=road['Moto'];road['BICI']=road['Bicicleta']
    road['campana']=road.Temporada.str.strip()
    point=d.loc[d['Lugar encuestado'].map(prev.norm).str.contains('cerro castillo')].copy()
    point['fecha_registro']=pd.to_datetime(point['Día']).dt.strftime('%Y-%m-%d')
    fpath=points
    orig=prev.cache_frame(fpath,'arcos_ptocontrol','factores_arcos_matrices',dpoints,False)
    orig=orig.loc[orig['Lugar encuestado'].map(prev.norm).str.contains('cerro castillo')]
    factors=[];terra=[];control_balance=[];audit=[]
    for (campaign,direction),g in road.groupby(['campana','SENTIDO']):
        days=g['DÍA'].nunique();expected_dates=set(cfg['calendario_Cerro_Castillo'][campaign])
        p=point.loc[point.campana.eq(campaign)&point['Sentido Viaje'].eq(direction)&point.fecha_registro.isin(expected_dates)]
        for day,gd in g.groupby('DÍA'):
            assert len(gd)==gd.CUARTO.nunique()==40
        for seg in ['VL','CAM','MOTO','BICI']:
            z=p.loc[p.segmento.eq(seg)];n=len(z);count=g[seg].sum(min_count=1);daily=count/days
            weight=daily/n if n else math.nan
            no=orig.loc[orig.Temporada.str.strip().eq(campaign)&orig['Sentido Viaje'].eq(direction)].iloc[0]
            orig_n=no['Suma de Encuesta VL' if seg=='VL' else 'Suma de Encuesta Cam'] if seg in ['VL','CAM'] else math.nan
            orig_q=no['Flujo Vehículos Livianos' if seg=='VL' else 'Flujo Camiones'] if seg in ['VL','CAM'] else math.nan
            orig_m=no['Factor de muestra vehiculos' if seg=='VL' else 'Factor de muestra camiones'] if seg in ['VL','CAM'] else math.nan
            exact=n and n==orig_n and math.isclose(daily,orig_q,rel_tol=1e-9)
            if exact:
                weight=1/orig_m
                assert math.isclose(weight,daily/n,rel_tol=1e-9)
            factors.append(dict(campana=campaign,estacion=cfg['estacion_campana'][campaign],sentido_control=direction,segmento=seg,
                conteos_acumulados=count,jornadas=days,horas_por_jornada=len(g)/days/4,conteo_medio_por_jornada=daily,
                encuestas_validas_compatibles=n,encuestas_numerador_ARISTO=orig_n,flujo_cache_ARISTO=orig_q,
                factor_muestra_ARISTO=orig_m,factor_expansion_ARISTO=1/orig_m if pd.notna(orig_m) and orig_m else math.nan,
                factor_expansion_utilizado=weight,reutiliza_factor_ARISTO_exacto=bool(exact),
                motivo='misma_clase_y_numerador_flujo' if exact else 'reconciliacion_clase_fisica_y_version_conteo; n_sin_pasajeros_bus_van; VL_sin_motos_bicis',
                unidad_factor='vehiculos_por_dia_de_jornada_10h_por_encuesta'))
            control_balance.append(dict(campana=campaign,sentido_control=direction,segmento=seg,conteo_medio_jornada=daily,
                encuestas_validas=n,suma_expansion_todo_control=weight*n if n else math.nan,
                flujo_control_sin_muestra=daily if not n else 0,NO_equivale_total_OD_proyecto=True))
            for route,zz in z.groupby('ruta'):
                if route!='terrestre_chile':continue
                # Solo OD nucleo + otras relaciones norte-sur previamente observadas en el cruce.
                allowed={(o,t) for o,t in zip(lake['Nombre origen'],lake['Nombre destino']) if 'Chile Chico' in (o,t) and o!=t}
                zz=zz.loc[zz.nucleo|pd.Series([(o,t) in allowed for o,t in zip(zz['Nombre origen'],zz['Nombre destino'])],index=zz.index)]
                for r in agrupar_encuestas(zz).itertuples():
                    terra.append(dict(**metadatos(r),ruta='terrestre_chile',grado_evidencia=r.grado_evidencia,encuestas_validas=r.encuestas_validas,
                        sentido_control=direction,n_estrato=n,factor_expansion=weight,
                        demanda_expandida_viajes_diarios_promedio=weight*r.encuestas_validas,demanda_diaria_24h_equivalente=math.nan,
                        cobertura_horas_dia=10,periodo_representado='jornada_diaria_media_campana_'+campaign,
                        estacion_completa=False,procedimiento='(conteo_clase/jornadas)/N_validos_control * n_OD_tierra_Chile',
                        fuente='ARISTO Carretera ODC3+BBDD1/BBDD2; arcos_ptocontrol cotejado',
                        validacion='expansion_parcial_condicional_jornada_10h; supuesto_uniformidad_seleccion_y_validez',
                        limitaciones='NO_24h; horarios_registro_no_exposicion_certificada; no_intervalos_confianza; no_otono_entero',
                        pasos_fronterizos=r.pasos_fronterizos))
        # Auditoria de fechas y horarios: no asumir Hora=hora efectiva de entrevista.
        start=int(g.CUARTO.min()//100);end=int(g.CUARTO.max()//100)+1
        h=p.Hora.map(lambda x: int(str(x).split(':')[0])+int(str(x).split(':')[1])/60 if pd.notna(x) else math.nan)
        audit.append(dict(campana=campaign,sentido_control=direction,fechas_registro=';'.join(sorted(p.fecha_registro.unique())),
            fechas_informe=';'.join(sorted(expected_dates)),coinciden_fechas=set(p.fecha_registro.unique())==expected_dates,
            encuestas=len(p),hora_inicio_conteo=start,hora_fin_conteo=end,registros_fuera_ventana_conteo=int(((h<start)|(h>=end)).sum()),
            nota='Hora puede ser creacion; no se filtra sin protocolo. Composicion de encuesta se supone representativa de10h para expansion parcial'))
    terra=pd.DataFrame(terra);save(terra,'matriz_od_terrestre')
    save(pd.DataFrame(factors),'factores_y_comparacion_ARISTO');save(pd.DataFrame(audit),'auditoria_exposicion_Cerro_Castillo')
    save(pd.DataFrame(control_balance),'balance_control_Cerro_NO_total_OD')
    # Los casos de Argentina/itinerario parcial no se agregan a Chile ni se expanden con fronteras.
    observed_pairs=set(zip(lake['Nombre origen'],lake['Nombre destino']))
    scope=d.loc[d.nucleo|d.ruta.eq('barcaza_actual')|pd.Series([(o,t) in observed_pairs and 'Chile Chico' in (o,t) for o,t in zip(d['Nombre origen'],d['Nombre destino'])],index=d.index)]
    classification=agrupar_encuestas(scope)
    units={'VL':'vehiculos_livianos','CAM':'vehiculos_carga','PEATON':'personas_peatones','MOTO':'motocicletas','BICI':'bicicletas',
           'OTRO_PASAJERO':'pasajero_individual_sin_unidad_fisica_asignable'}
    classification['unidad']=classification.segmento.map(units).fillna('unidad_fisica_no_clasificada')
    classification['unidad_encuestas']='registros_validos_de_encuesta'
    classification['posibilidad_expansion']=classification.ruta.map({'barcaza_actual':'solo_interceptadas_cruce',
        'terrestre_chile':'parcial_jornada10h_si_Cerro_Castillo','terrestre_argentina':'sin_marco_OD_circuito_fronteras',
        'terrestre_pais_desconocido':'no_asignable_a_Chile','ruta_desconocida':'no_asignable_por_ruta'})
    save(classification,'clasificacion_rutas_y_pasos')
    unknown=classification.loc[~classification.ruta.isin(['barcaza_actual','terrestre_chile'])].copy()
    unknown['demanda_expandida_viajes_diarios_promedio']=math.nan;unknown['factor_expansion']=math.nan
    unknown['demanda_diaria_24h_equivalente']=math.nan
    unknown['fuente']='ARISTO EOD; sin_ponderacion_OD_ruta_verificada'
    save(unknown,'matriz_od_desconocidos')
    arg=unknown.loc[unknown.ruta.eq('terrestre_argentina')].copy();save(arg,'observaciones_argentina_sin_expansion')
    # Incluir celdas no medidas como faltantes, nunca asignarles cero poblacional.
    placeholders=[]
    nucleus_geo=scope.loc[scope.nucleo,['origen','destino','nivel_origen','nivel_destino']].drop_duplicates()
    for r in nucleus_geo.itertuples():
        for station in ['Verano','Otoño','Invierno','Primavera']:
            for seg in ['VL','CAM','PEATON']:
                for route,frame in [('barcaza_actual',bar),('terrestre_chile',terra)]:
                    match=frame.estacion.eq(station)&frame.origen.eq(r.origen)&frame.destino.eq(r.destino)&frame.segmento.eq(seg)
                    if not match.any():
                        placeholders.append(dict(estacion=station,origen=r.origen,destino=r.destino,sentido=r.origen+' → '+r.destino,
                            segmento=seg,ruta=route,nivel_origen=r.nivel_origen,nivel_destino=r.nivel_destino,ambito='nucleo_Chile_Chico_ribera_norte',
                            encuestas_validas=0,demanda_expandida_viajes_diarios_promedio=math.nan,demanda_diaria_24h_equivalente=math.nan,factor_expansion=math.nan,
                            validacion='sin_medicion_o_cero_muestral_no_es_cero_poblacional',limitaciones='no_completar_con_cero',
                            periodo_representado='sin_estimacion',estacion_completa=False))
    ph=pd.DataFrame(placeholders)
    bar=pd.concat([bar,ph.loc[ph.ruta.eq('barcaza_actual')]],ignore_index=True)
    terra=pd.concat([terra,ph.loc[ph.ruta.eq('terrestre_chile')]],ignore_index=True)
    save(bar,'matriz_od_barcaza');save(terra,'matriz_od_terrestre')
    consolidated=pd.concat([bar,terra,unknown],ignore_index=True,sort=False)
    save(consolidated,'matriz_od_consolidada')
    # Composicion mensual heredada y geografia: trazabilidad, sin datos personales.
    save(refs,'flujos_base_mensuales_reutilizados')
    res=scope.groupby(['Nombre origen','Nombre destino','origen','destino','nivel_origen','nivel_destino']).size().reset_index(name='encuestas')
    save(res,'resolucion_territorial_sin_coordenadas_personales')
    source_rows=[
        ('Cap 6/BBDD EOD/Base de datos agregada.xlsx','BBDD-1; BBDD-2 (solo válidas)','Microdatos filtrados y clasificación de ruta/OD/segmento'),
        ('Cap 6/BBDD EOD/Base de datos agregada.xlsx','TD Transb, filas 20–25','Fórmulas originales y valores almacenados comparados con reconstrucción por conteo'),
        ('Cap 6/Mediciones de flujo/Consolidado Flujos.xlsx','Transbordadores; Carretera','Conteos de transbordador y ODC-3 Cerro Castillo'),
        ('Cap 8/EOD_Puntos_Pasada_Final.xlsx','arcos_ptocontrol','Revisión y cotejo de factores agregados por arco'),
        ('Cap 8/EOD_Puntos_Pasada_Final.xlsx','Puntos por OD (Ordenados)','Cobertura de rutas/puntos disponibles para enlace individual'),
        ('Cap 8/BBDD_para_Pijas_actualizada_EOD_Final.xlsx','Arcos por OD','Cobertura de arcos y pasos fronterizos codificados'),
        ('Cap 7/HDM4/Costos_Red_Vial_Actual.xlsx','Costos_Red_Vial; Tipos de arcos','Inventario de longitud, velocidades, costos operacionales y costo generalizado por arco'),
        ('Cap 10/Resumen_Rutas_Red_Vial_Actual.xlsx','Rutas Red Actual','Cobertura de alternativas de red por país y disponibilidad de OD'),
        ('Informe ARISTO.pdf','§2.2.6; Cap. 6; §8.1–8.2','Método/fuentes de campañas y expansión de matrices'),
        ('Estudios de consultoras/Informe 2c - MOP aysen vf.pdf','§6.1, pp. 75–90','Contraste muestral CITRA, no expansión'),
        ('INFORME_2_EVALSO-1.pdf','Metodología de demanda, versión local','Estructura futura de alternativas; sin logit en esta etapa')]
    manifest=[]
    for rel,sheets,use in source_rows:
        f=ROOT.parent/rel
        manifest.append(dict(archivo=rel,hoja_seccion_paginas=sheets,uso=use,sha256=hashlib.sha256(f.read_bytes()).hexdigest()))
    save(pd.DataFrame(manifest),'fuentes_utilizadas')
    save(coded,'cobertura_itinerarios_microdatos')
    citra=pd.DataFrame(citra_cfg['observaciones'],columns=['origen','destino','segmento','encuestas_CITRA','unidad','tabla'])
    citra['estacion']='Verano_CITRA'
    citra['ruta']='no_desagregada_en_CITRA'
    citra['flujo_expandido']=math.nan
    citra['fuente']='Informe 2 corregido CITRA; transcripcion directa de cuadros OD; no expandido'
    save(citra,'contraste_CITRA_matriz_muestral')
    route_book=ROOT.parent/'Cap 10/Resumen_Rutas_Red_Vial_Actual.xlsx'
    routes=pd.read_excel(route_book,sheet_name='Rutas Red Actual')
    zone_pairs=[(39,31),(31,39),(39,40),(40,39),(40,31),(31,40)]
    network_inventory=[]
    for country,n in routes['País_Ruta'].value_counts(dropna=False).items():
        network_inventory.append(dict(archivo='Cap 10/Resumen_Rutas_Red_Vial_Actual.xlsx',hoja='Rutas Red Actual',atributos='RUTA; Origen; Destino; País_Ruta; N_arcos; CostGen_VL; CO_VL; Tiempo_VL_h; Tiempo_VP_h',
            cobertura=f'{country}: {int(n)} rutas en matriz de red',OD_local_objetivo='no',n_rutas_locales_objetivo=0,
            nota='Inventario de alternativas de red agregadas; no es matriz expandida de viajes'))
    for o,t in zone_pairs:
        n=int(((routes.Origen==o)&(routes.Destino==t)).sum())
        network_inventory.append(dict(archivo='Cap 10/Resumen_Rutas_Red_Vial_Actual.xlsx',hoja='Rutas Red Actual',atributos='Origen/Destino son zonas; códigos 31 Coyhaique, 39 Chile Chico, 40 Río Ibáñez',
            cobertura='pares locales comprobados',OD_local_objetivo=f'{o}->{t}',n_rutas_locales_objetivo=n,
            nota='No aparece esta dirección OD en Resumen Rutas Red Actual'))
    save(pd.DataFrame(network_inventory),'inventario_atributos_rutas_y_cobertura')
    (OUT/'config_utilizada.json').write_text(json.dumps(cfg,ensure_ascii=False,indent=2))
    checks={'conservacion_barcaza':bool(all(math.isclose(r.referencia_dia,(r.asignado_OD if pd.notna(r.asignado_OD) else 0)+(r.sin_asignar_dia if pd.notna(r.sin_asignar_dia) else 0),rel_tol=1e-9,abs_tol=1e-9) for r in bal.dropna(subset=['referencia_dia']).itertuples())),
        'sin_flujos_argentina_en_chile':bool(terra.ruta.eq('terrestre_chile').all()),'sin_24h_tierra_inventado':bool(terra.demanda_diaria_24h_equivalente.isna().all()),
        'sin_ids_personales':bool(not (set(consolidated.columns)&{'ID 2','Patente','Latitud','Longitud'})),
        'cuatro_estaciones':bool(set(consolidated.estacion)=={'Verano','Otoño','Invierno','Primavera'})}
    assert all(checks.values()),checks
    (OUT/'validacion.json').write_text(json.dumps(checks,indent=2))
    summary(bar,terra,unknown,bal,pd.DataFrame(factors),pd.DataFrame(audit),classification,cfg,coded,pd.DataFrame(network_inventory),citra,tdcompare)
    print('Matrices principales y resumen generados en',OUT,flush=True)


def summary(bar,terra,unknown,balance,factors,audit,classes,cfg,coded,network_inventory,citra,tdcompare):
    # Mostrar celdas medidas en tabla corta; los faltantes completos estan en CSV.
    b=bar.loc[bar.ambito.eq('nucleo_Chile_Chico_ribera_norte')&bar.segmento.isin(['VL','CAM','PEATON'])&bar.demanda_expandida_viajes_diarios_promedio.notna()]
    t=terra.loc[terra.ambito.eq('nucleo_Chile_Chico_ribera_norte')&terra.demanda_expandida_viajes_diarios_promedio.notna()]
    cols=['estacion','origen','destino','segmento','encuestas_validas','factor_expansion','demanda_expandida_viajes_diarios_promedio']
    core='nucleo_Chile_Chico_ribera_norte'
    cls=classes.loc[classes.ambito.eq(core)].copy()
    keys=['estacion','origen','destino','sentido','segmento','ruta']
    cls=cls.groupby(keys,dropna=False).agg(
        grado_evidencia=('grado_evidencia',lambda s:'; '.join(sorted(set(s.dropna().astype(str))))),
        pasos_fronterizos=('pasos_fronterizos',lambda s:'; '.join(sorted(set(s.dropna().astype(str))))),
        encuestas_validas=('encuestas_validas','sum'),
        posibilidad_expansion=('posibilidad_expansion',lambda s:'; '.join(sorted(set(s.dropna().astype(str)))))).reset_index()
    units={'VL':'vehiculos_livianos','CAM':'vehiculos_carga','PEATON':'personas_peatones','MOTO':'motocicletas','BICI':'bicicletas',
           'OTRO_PASAJERO':'pasajero_individual_sin_unidad_fisica_asignable'}
    cls['unidad']=cls.segmento.map(units).fillna('unidad_fisica_no_clasificada')
    cls['demanda_expandida_viajes_diarios_promedio']=float('nan')
    for name,frame in [('barcaza_actual',bar),('terrestre_chile',terra)]:
        vals=frame.groupby(keys,dropna=False).demanda_expandida_viajes_diarios_promedio.sum(min_count=1).reset_index()
        vals=vals.loc[vals.ruta.eq(name),keys+['demanda_expandida_viajes_diarios_promedio']]
        cls=cls.merge(vals,on=keys,how='left',suffixes=('','_expanded'))
        col='demanda_expandida_viajes_diarios_promedio_expanded'
        cls['demanda_expandida_viajes_diarios_promedio']=cls[col].combine_first(cls['demanda_expandida_viajes_diarios_promedio'])
        cls=cls.drop(columns=col)
    cls=cls.sort_values(['estacion','origen','destino','segmento','ruta'])
    totals=[]
    for name,frame in [('barcaza_actual',bar),('terrestre_chile',terra)]:
        q=frame.loc[frame.ambito.eq(core)].groupby(['estacion','sentido','segmento','unidad'],dropna=False).agg(
            encuestas_validas=('encuestas_validas','sum'),
            demanda_expandida_viajes_diarios_promedio=('demanda_expandida_viajes_diarios_promedio',lambda s:s.sum(min_count=1))).reset_index()
        q['ruta']=name; totals.append(q)
    total_table=pd.concat(totals,ignore_index=True)
    total_table=total_table.loc[total_table.encuestas_validas.gt(0)|total_table.demanda_expandida_viajes_diarios_promedio.notna()]
    unk=unknown.loc[unknown.ambito.eq(core)].groupby(['estacion','origen','destino','segmento','ruta','grado_evidencia','pasos_fronterizos'],dropna=False).encuestas_validas.sum().reset_index()
    unk['demanda_expandida_viajes_diarios_promedio']=float('nan')
    txt=f'''# Matrices OD de referencia por ruta — 25/10/2026

Esta entrega estima barcaza y, de forma parcial/condicional, tierra por Chile. No ejecuta SBO, ocupacion, captacion ni logit. Alternativas posteriores: tierra Chile, barcaza actual y tierra Argentina, sujetas a factibilidad por OD/segmento; con proyecto se agrega Bahia Jara–Levican. No se declaran tres alternativas disponibles para todos los OD solo por incluirlas en la estructura.

## Matriz expandida de barcaza: nucleo local

Unidades: VL vehiculos de cuatro ruedas; CAM agregado de carga ARISTO (incluye furgones/remolques); PEATON personas a pie, no todos los pasajeros. MOTO y BICI estan separados en CSV. Se evita convertir respuestas de pasajeros de bus/van/transfer en vehiculos. El tipo generico Vehiculo privado de Normal se aproxima a VL: requiere validar su composicion interna.

Frecuencia ordinaria11 diciembre–abril y8 mayo–noviembre zarpes semanales por sentido, excluida carga peligrosa (usuario). Conteos no modificados. Q_base=(media_por_fila_de_conteo×f/7); Q_OD=Q_base×n_OD/n_estrato. El script reconstruye la media directamente desde `Consolidado Flujos.xlsx/Transbordadores`, sin ejecutar el módulo SBO. Primavera tiene apoyo de tabla78 para dos zarpes por sentido; asumir una fila por zarpe en verano/mayo sigue condicionado. La hoja no tiene una marca que permita verificar si un conteo corresponde a un servicio de carga peligrosa; si existiera una fila de ese tipo dentro de las campañas, no puede separarse con lo disponible. Composicion de encuestas se supone representativa; para peatones no se multiplican grupos ni ocupantes.

`TD Transb` original usa `G=F×D/7` y `H=F×E/7` con F=9 y conserva sus valores calculados. En primavera, D/E corresponden a la suma de dos filas de conteo por sentido, mientras §2.2.6 describe media por zarpe multiplicada por zarpes semanales. Por eso se muestra intacto el resultado de ARISTO y se reconstruye por media de filas y frecuencia ordinaria confirmada; la reconstrucción no modifica ni reemplaza el libro. La separación de bicicletas y motos desde “vehículos privados” evita mezclarlas con VL de cuatro ruedas.

{prev.mdtable(tdcompare[['fila_TD_Transb','campana','sentido','formula_G','valor_G_cache','formula_H','valor_H_cache','filas_conteo_observadas','conteo_D_coincide_sum_filas','conteo_E_coincide_sum_filas','frecuencia_ordinaria_actual','flujo_privados_actual','flujo_CAM_actual']])}

Verano y primavera son transferencia de perfil de campaña a la estacion; **Otoño en esta tabla significa solo mayo**, no marzo/abril. Invierno y celdas sin muestra quedan NaN. No se hace anualizacion. `estacion_completa` y `periodo_representado` indican alcance.

{prev.mdtable(b[cols].sort_values(['segmento','estacion','origen','destino']))}

## Matriz terrestre por Chile: expansion parcial de jornadas contadas

Las filas siguientes son **vehiculos por dia promedio dentro de las10horas contadas**, NO viajes/dia de24h. La columna `demanda_diaria_24h_equivalente` queda NaN. Son resultados condicionales con datos reales de conteo y encuesta, no flujos poblacionales diarios certificados ni factores horarios inventados.

{prev.mdtable(t[cols].sort_values(['segmento','estacion','origen','destino']))}

## Cobertura observada por ruta y grado de identificación (OD núcleo)

Esta tabla distingue la matriz expandida de la composición de la muestra. En rutas Argentina, territorio desconocido y ruta desconocida, `demanda_expandida_viajes_diarios_promedio` queda faltante: los conteos de pasos no son factores OD.

De las **27 encuestas previamente consideradas compatibles con “sin transbordador”** en los cuatro OD núcleo, la verificación adicional de campos fronterizos reclasifica 20 como terrestre por Chile (declaración de ningún paso, control de carretera en Chile) y 7 como ruta argentina probable (pasos declarados, pero sin trazado completo enlazado). No son 27 viajes terrestres por Chile. Los siete casos argentinos no se expanden.

{prev.mdtable(cls)}

### Totales del núcleo por estación, sentido y segmento

Los totales terrestres representan solo la jornada de 10 horas y los OD clasificados; los de barcaza son promedios diarios equivalentes según frecuencia ordinaria. No representan un mercado completo y no se suman como si tuvieran cobertura temporal idéntica.

{prev.mdtable(total_table.sort_values(['estacion','sentido','segmento','ruta']))}

### Encuestas cuyo flujo poblacional por ruta sigue sin identificar

Se informa el número de encuestas, pero no se rellena la demanda expandida con cero. Las rutas por Argentina pueden estar declaradas o ser probables según los pasos informados.

{prev.mdtable(unk)}

Se usa solamente ODC-3 Cerro Castillo: F=(C_acumulado/jornadas)/N_validos_control; Q_OD_CH=F×n_OD_CH. Numerador del OD exige negativa explicita de transbordadores Y de pasos fronterizos, sin otro campo positivo/contradictorio y sin interceptacion extranjera. No basta Ninguno en transbordadores. El denominador contiene TODOS los validos de la misma campaña/sentido/clase y fechas, no solo los OD del proyecto ni solo quienes declararon tierra. Asi no se asigna todo Cerro Castillo a Chile Chico–Coyhaique.

No se suman otros controles ni fronteras. No se expanden declaraciones de barcaza fuera del cruce como si fueran demanda adicional: los pesos de carretera no se agregan a la matriz lacustre. Las rutas desconocidas no se completan con Chile. Los ceros de muestra y estratos sin respuestas permanecen sin estimacion poblacional.

La identidad F=flujo/encuestas corresponde a la inversa del cociente J/K de arcos_ptocontrol, pero SOLO se reutiliza literalmente cuando coinciden clases, numerador y flujo. CAM conserva el factor original donde coincide. VL se reconcilia a cuatro ruedas, excluye respuestas de pasajeros y usa los conteos de la version Cap6: la comparación preserva los valores originales de Cap8, que difieren en algunos flujos de Cerro Castillo. No se corrigio ningun original ni se adopto un factor mensual como diario.

{prev.mdtable(factors.loc[factors.segmento.isin(['VL','CAM']),['campana','sentido_control','segmento','conteo_medio_por_jornada','encuestas_validas_compatibles','flujo_cache_ARISTO','encuestas_numerador_ARISTO','factor_expansion_utilizado','reutiliza_factor_ARISTO_exacto']])}

## Exposicion, seleccion y alcance de la estimacion terrestre

Conteos40cuartos/jornada=10h; fechas EOD cotejadas con ARISTO tabla75 primavera, planificacion verano y tabla140 mayo. La hoja de conteos solo etiqueta Dia1/Dia2; el enlace de fechas se hace por campaña/documento, no por ID de jornada enlazado individualmente.

{prev.mdtable(audit)}

Hora de registro puede ser creacion, no entrevista; existen diferencias con la ventana de conteo que no se resolvieron filtrando arbitrariamente. **Supuesto de esta expansion parcial:** composicion de validos representa la jornada, con seleccion/no respuesta/validez ignorables dentro del estrato. No es un ponderador de inclusion individual certificado. Para un flujo24h y representatividad probabilistica faltan protocolo de seleccion y no respuesta, hora efectiva y cobertura de encuesta, factor de jornada a24h y calendario estacional. Estos limites no impiden mostrar el calculo parcial, pero impiden presentarlo como el total diario del mercado. No se generaron intervalos de confianza de un muestreo cuyo diseño no esta disponible.

## Conservacion y geografia

`conservacion_flujos_barcaza.csv` comprueba por estacion/sentido/clase: asignado_OD+sin_asignar=flujo_base. `flujos_barcaza_sin_OD.csv` conserva conteos positivos sin muestra compatible y componentes ausentes. Toda la muestra del cruce se conserva para balance; `ambito` separa el nucleo local de otros OD realmente observados en el cruce cuya elegibilidad para el proyecto NO se confirma automaticamente. No se incorporan todos los OD de Aysen.

`balance_control_Cerro_NO_total_OD.csv` conserva la identidad entre peso×muestra completa y conteo por jornada. Ese balance de control NO es demanda total de los OD del proyecto. No se suman unidades de personas y vehiculos, ni jornada10h con equivalentes24h como si tuvieran la misma cobertura.

Rio Ibanez se conserva como comuna/zona salvo mencion textual UNICA de una localidad en Busqueda/Direccion EOD. Esta etapa no usa coordenadas ni publica domicilios; no define localidad por radio arbitrario. La resolucion y escala quedan en columnas `nivel_origen`, `nivel_destino` y en `resolucion_territorial_sin_coordenadas_personales.csv`. Coyhaique y Chile Chico conservan escala BBDD2 si no hay delimitacion de localidad comprobada. No se asigna todo Rio Ibanez a Puerto Ibanez. Una mencion textual no es validacion de limites urbanos.

## Argentina y desconocidos: complemento, sin sumar a Chile

`clasificacion_rutas_y_pasos.csv` entrega estacion, OD, segmento, ruta, pasos identificados, encuestas, evidencia y posibilidad de expansion. Tierra con pais desconocido queda separada. Uno o varios pasos argentinos declarados con ningún transbordador se etiquetan como ruta argentina probable, no como itinerario completo confirmado. Un control fronterizo por sí solo no prueba el recorrido completo. Los campos dicen usados/por usar y no confirman ejecución efectiva de todos los tramos. No se asigna el conteo completo fronterizo a un OD ni se suman dos pasos del mismo circuito. Argentina sigue sin expansión OD verificada.

### Itinerarios codificados y atributos de costos

Se enlazaron `Puntos por OD (Ordenados)` y `Arcos por OD` por identificador solo para comprobar cobertura, sin exponer registros individuales. Para el núcleo local no hay rutas, nodos ni arcos codificados enlazables; esto impide validar el trazado completo OD desde esas hojas. En `Cap 10/Resumen_Rutas_Red_Vial_Actual.xlsx`, la hoja `Rutas Red Actual` informa país, número de arcos, `CostGen_VL`, `CO_VL` y tiempos VL/VP; incluye alternativas por Chile y Argentina en la red. Sin embargo, no contiene rutas para los pares zonales locales comprobados (31 Coyhaique, 39 Chile Chico y 40 Río Ibáñez). `Cap 7/HDM4/Costos_Red_Vial_Actual.xlsx`, hoja `Costos_Red_Vial`, ofrece por arco rutas, longitud, velocidades VL/VP, costos operacionales y costo generalizado VL; `Tipos de arcos` agrega atributos viales. No basta para calcular ahora el costo generalizado argentino del núcleo: faltan trayectos conectados del OD local, verificación del paso/ruta, componentes de cruce fronterizo y compatibilidad de tiempo/costo por clase.

{prev.mdtable(network_inventory)}

Cobertura de rutas codificadas de microdatos para OD núcleo:

{prev.mdtable(coded.groupby('estacion',dropna=False).agg(encuestas_validas=('encuestas_validas','sum'),registros_con_ruta_codificada=('registros_con_ruta_codificada','sum'),registros_con_nodos_ruta=('registros_con_nodos_ruta','sum'),registros_con_arcos=('registros_con_arcos','sum')).reset_index())}

`matriz_od_desconocidos.csv` contiene rutas desconocidas, pais terrestre no identificado y observaciones argentinas sin flujo verificable; cada ruta queda rotulada por separado. `observaciones_argentina_sin_expansion.csv` permite examinarlas sin mezclarlas con Chile. Si una declaracion de barcaza fuera del cruce no acredita uso efectivo, conserva esa evidencia auxiliar sin duplicar el flujo ya expandido del cruce.

### Contraste independiente CITRA

CITRA publicó estas cantidades de encuestas OD válidas por pares en su campaña estival del 20 de enero al 5 de febrero de 2026. Son registros muestrales, agregados sobre rutas no desagregadas, no flujos expandidos y no se suman a ARISTO. El cuadro 2 (p. 83) muestra 164 viajes Coyhaique→Chile Chico VL; el archivo derivado previo `outputs/sbo_actualizada/E_contraste_CITRA_muestras_PDF.csv` indicaba 82. Esta discrepancia se resuelve contra el PDF original. No se asigna ninguna fila de CITRA a una alternativa de ruta.

{prev.mdtable(citra[['origen','destino','segmento','encuestas_CITRA','unidad','tabla','ruta','flujo_expandido']])}

## Archivos y continuacion

Matrices principales: matriz_od_barcaza.csv, matriz_od_terrestre.csv, matriz_od_desconocidos.csv y matriz_od_consolidada.csv. `demanda_expandida_viajes_diarios_promedio` requiere leer SIEMPRE cobertura/periodo; la columna24h distingue lo que no puede calcularse en tierra. No se totalizan rutas para declarar mercado completo mientras falten cobertura y rutas.

Reproduccion: `python scripts/07_matrices_od_referencia.py`. Configuracion: data/templates/config_matrices_od.json. El script crea o reusa caches locales seleccionados y versionados por hash; no requiere correr el script03 de estimacion antigua ni el módulo SBO. Fuentes: ARISTO EOD/Consolidado/arcos_ptocontrol; originales intactos, sin commit/push. Datos intermedios geográficos y claves permanecen en data/interim ignorado por Git. CITRA sigue solo como contraste de cuadros PDF sin expansión ni combinación de muestras. Se preservan los módulos SBO para retomarlos. Los costos generalizados y el logit quedan pendientes.
'''
    (OUT/'resumen_matrices.md').write_text(documento_legible(txt))


if __name__=='__main__':main()
